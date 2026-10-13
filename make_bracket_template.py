# /// script
# requires-python = ">=3.12"
# dependencies = ["python-docx~=1.1"]
# ///
"""Generate templates/bracket.docx, the Word twin of themes/bracket.css.

Beyond the style ids every template has, it opts into build_cv.py's optional parts: Photo, Contact and
Highlights placeholders, skill bars, a role timeline, tool chips and quote bars.
"""
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_UNDERLINE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

from make_pro_template import style

FONT = "Montserrat"
INK, MUTED, ACCENT, SIDE, TRACK, LINE = "262626", "7A7A7A", "F6B92B", "F4F3EF", "DDDBD3", "D9D9D9"


def font(**kw):
    return style(font=FONT, **kw)


# Schema order of w:pPr children that may follow the ones added here; Word rejects out-of-order elements.
AFTER_NUMPR = ("w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
               "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
               "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap",
               "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
               "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def ppr(s, xml: str, after: tuple[str, ...]) -> None:
    """Add a w:pPr child, given as XML without namespace declarations, before its schema successors."""
    root = xml.index(">")
    s.element.get_or_add_pPr().insert_element_before(parse_xml(f"{xml[:root]} {nsdecls('w')}{xml[root:]}"), *after)


def border(s, sides: str, color: str, size: int, space: int = 4) -> None:
    """sides: space-separated w:pBdr sides. Adjacent paragraphs with equal borders share one box, so ruled
    lists need "top between bottom"."""
    order = ("top", "left", "bottom", "right", "between")  # schema order
    edges = "".join(f'<w:{side} w:val="single" w:sz="{size}" w:space="{space}" w:color="{color}"/>'
                    for side in sorted(sides.split(), key=order.index))
    ppr(s, f"<w:pBdr>{edges}</w:pBdr>", AFTER_NUMPR[2:])


def char_style(doc, name: str, *, color: str | None = None, fill: str | None = None, size: float | None = None, bold=None):
    s = doc.styles.add_style(name, WD_STYLE_TYPE.CHARACTER)
    s.font.name, s.font.bold = FONT, bold
    if color:
        s.font.color.rgb = RGBColor.from_string(color)
    if size:
        s.font.size = Pt(size)
    if fill:
        s.element.get_or_add_rPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>'))
    return s


def marks(doc) -> dict[str, str]:
    """Bullet list first (build_cv.py takes the first bullet numbering), then the diamond and ring marks."""
    numbering = doc.part.numbering_part.element
    defs = {"list": ("•", ACCENT, 284, 227, 9), "diamond": ("◆", ACCENT, 340, 340, 8), "ring": ("○", INK, 340, 340, 9)}
    for i, (key, (glyph, color, left, hanging, size)) in enumerate(defs.items()):
        n = 90 + i
        numbering.insert(i, parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{n}"><w:lvl w:ilvl="0"><w:start w:val="1"/>'
            f'<w:numFmt w:val="bullet"/><w:lvlText w:val="{glyph}"/><w:lvlJc w:val="left"/>'
            f'<w:pPr><w:ind w:left="{left}" w:hanging="{hanging}"/></w:pPr>'
            f'<w:rPr><w:rFonts w:ascii="Segoe UI Symbol" w:hAnsi="Segoe UI Symbol"/><w:color w:val="{color}"/>'
            f'<w:sz w:val="{size * 2}"/></w:rPr></w:lvl></w:abstractNum>'))
        numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="{n}"><w:abstractNumId w:val="{n}"/></w:num>'))
    return {key: str(90 + i) for i, key in enumerate(defs)}


def numbered(s, num_id: str) -> None:
    ppr(s, f'<w:numPr><w:numId w:val="{num_id}"/></w:numPr>', AFTER_NUMPR)


def cell(c, width_cm: float, *, fill: str | None = None, pad=(0, 0, 0, 0), borders: dict | None = None, valign=None) -> None:
    """pad: top, right, bottom, left in cm; borders: side -> (color, eighths of a point)."""
    c.width = Cm(width_cm)
    tc_pr = c._tc.get_or_add_tcPr()
    if borders:
        order = ("top", "left", "bottom", "right")  # schema order
        tc_pr.append(parse_xml(f'<w:tcBorders {nsdecls("w")}>' + "".join(
            f'<w:{side} w:val="single" w:sz="{borders[side][1]}" w:space="0" w:color="{borders[side][0]}"/>'
            for side in order if side in borders) + "</w:tcBorders>"))
    if fill:
        tc_pr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>'))
    margin = dict(zip(("top", "right", "bottom", "left"), pad))
    tc_pr.append(parse_xml(f'<w:tcMar {nsdecls("w")}>' + "".join(
        f'<w:{side} w:w="{int(margin[side] * 567)}" w:type="dxa"/>' for side in ("top", "left", "bottom", "right"))  # schema order
        + "</w:tcMar>"))
    if valign:
        c.vertical_alignment = valign


def table(container, widths: list[float], rows: int = 1):
    t = container.add_table(rows=rows, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    for col, w in zip(t._tbl.tblGrid.gridCol_lst, widths):
        col.w = Cm(w)
    t._tbl.tblPr.get_or_add_tblLayout().set(qn("w:type"), "fixed")
    t._tbl.tblPr.insert_element_before(parse_xml(
        f'<w:tblCellMar {nsdecls("w")}><w:left w:w="0" w:type="dxa"/><w:right w:w="0" w:type="dxa"/></w:tblCellMar>'),
        "w:tblLook", "w:tblCaption", "w:tblDescription")
    return t


def main() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(0)
    sec.top_margin, sec.bottom_margin = Cm(1.2), Cm(1.4)
    sec.header_distance, sec.footer_distance = Cm(0), Cm(0)
    ids = marks(doc)

    font(doc=doc, name="Normal", size=9, after=4, line=1.3)
    font(doc=doc, name="Name", base="Normal", size=27, bold=True, caps=True, after=0, line=1.0)
    font(doc=doc, name="Job Title", base="Normal", size=7, color=MUTED, caps=True, spacing=1.3, after=3)
    font(doc=doc, name="CV Photo", base="Normal", size=2, after=0, line=1.0)
    contact = font(doc=doc, name="CV Contact", base="Normal", size=8, bold=True, before=0, after=0, line=1.0)
    contact.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    border(contact, "top between bottom", INK, 6, 5)
    font(doc=doc, name="CV Highlights", base="Normal", size=2, after=0)
    value = font(doc=doc, name="CV Highlight Value", base="Normal", size=18, bold=True, after=3, line=1.0)
    value.font.underline = WD_UNDERLINE.THICK
    value.element.rPr.find(qn("w:u")).set(qn("w:color"), ACCENT)
    font(doc=doc, name="CV Highlight Label", base="Normal", size=7.5, color=MUTED, after=0)

    h1 = font(doc=doc, name="Heading 1", base="Normal", size=13, bold=True, before=0, after=8)
    numbered(h1, ids["diamond"])
    font(doc=doc, name="Heading 2", base="Normal", size=9.5, bold=True, before=4, after=0)
    role = font(doc=doc, name="CV Role", base="Heading 2", size=9.5, bold=True, before=8, after=0)
    numbered(role, ids["ring"])
    company = font(doc=doc, name="CV Company", base="Normal", size=8.2, color=MUTED, after=1)
    company.paragraph_format.left_indent = Cm(0.6)
    date = font(doc=doc, name="Date", base="Normal", size=7.2, bold=True, spacing=0.3, after=5)
    date.font.underline = WD_UNDERLINE.THICK
    date.element.rPr.find(qn("w:u")).set(qn("w:color"), ACCENT)
    h3 = font(doc=doc, name="Heading 3", base="Normal", size=10.5, bold=True, before=0, after=5)
    numbered(h3, ids["diamond"])
    font(doc=doc, name="Heading 4", base="Normal", size=6.8, color=MUTED, bold=True, caps=True, spacing=1.1, before=5, after=0)
    font(doc=doc, name="Sidebar Text", base="Normal", size=8.5, after=1)
    kv = font(doc=doc, name="CV Key Value", base="Normal", size=8.5, after=3)
    kv.paragraph_format.tab_stops.add_tab_stop(Cm(4.6), alignment=2)  # right, at the sidebar text edge
    font(doc=doc, name="CV Skill", base="Normal", size=8.5, after=2)
    font(doc=doc, name="CV Skill Bar", base="Normal", size=1, after=5, line=1.0)
    font(doc=doc, name="List Paragraph", base="Normal", size=9, after=2)
    tools = font(doc=doc, name="CV Tools", base="Normal", size=7, after=8, line=1.7)
    tools.paragraph_format.left_indent = Cm(0.6)
    quote = font(doc=doc, name="CV Quote", base="Normal", size=9, color="3A3A3A", after=3)
    quote.font.italic = True
    quote.paragraph_format.left_indent = Cm(0.4)
    border(quote, "left", ACCENT, 18, 8)
    font(doc=doc, name="Sidebar Sections Spacing", base="Normal", size=4, after=14)
    font(doc=doc, name="Main Sections Spacing", base="Normal", size=4, after=22)
    char_style(doc, "CV Accent", color=ACCENT)
    char_style(doc, "CV Muted", color=MUTED, size=7.5)
    char_style(doc, "CV Tool", fill=SIDE, size=7)
    char_style(doc, "CV Skill Fill", fill=ACCENT)
    char_style(doc, "CV Skill Track", fill=TRACK)
    for s in ("CV Contact", "CV Role", "CV Company", "Date", "Heading 1", "Heading 3"):
        doc.styles[s].paragraph_format.keep_with_next = True

    # Header: [bracket around photo | role over name | contact], headline numbers below.
    head = table(doc, [4.7, 9.6, 6.0], rows=2)
    head._tbl.tblPr.insert_element_before(parse_xml(f'<w:tblInd {nsdecls("w")} w:w="{int(0.7 * 567)}" w:type="dxa"/>'),  # bracket 7 mm in
                                          "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
    photo_c, name_c, contact_c = head.rows[0].cells
    cell(photo_c, 4.7, pad=(0.5, 0, 0.3, 0.6), borders={"top": (ACCENT, 96), "left": (ACCENT, 96)},
         valign=WD_CELL_VERTICAL_ALIGNMENT.CENTER)
    cell(name_c, 9.6, pad=(0.9, 0.4, 0.3, 0.2), valign=WD_CELL_VERTICAL_ALIGNMENT.CENTER)
    cell(contact_c, 6.0, pad=(0.9, 1.4, 0.3, 0.4), valign=WD_CELL_VERTICAL_ALIGNMENT.CENTER)
    frame = table(photo_c, [3.4])  # the photo's yellow offset block
    cell(frame.rows[0].cells[0], 3.4, borders={"right": (ACCENT, 96), "bottom": (ACCENT, 96)})
    photo_c._tc.remove(photo_c.paragraphs[0]._p)
    photo_c.paragraphs[-1]._p.style = "CVPhoto"
    frame.rows[0].cells[0].paragraphs[0]._p.style = "CVPhoto"
    frame.rows[0].cells[0].paragraphs[0].add_run("{Photo 3.1 gray}")
    name_c.paragraphs[0]._p.style = "JobTitle"
    name_c.paragraphs[0].add_run("{Job Title}")
    name_c.add_paragraph(style="Name").add_run("{Name}")
    contact_c.paragraphs[0]._p.style = "CVContact"
    contact_c.paragraphs[0].add_run("{Contact}")
    stats = head.rows[1].cells[0].merge(head.rows[1].cells[2])
    cell(stats, 20.3, pad=(0.6, 1.4, 0.8, 0.9))
    stats.paragraphs[0]._p.style = "CVHighlights"
    stats.paragraphs[0].add_run("{Highlights}")

    # Body: light sidebar, main column.
    doc.add_paragraph(style="CV Photo")  # keeps Word from merging the two tables
    lay = table(doc, [6.4, 14.6])
    side, main_c = lay.rows[0].cells
    cell(side, 6.4, fill=SIDE, pad=(1.0, 0.8, 1.2, 1.0))
    cell(main_c, 14.6, pad=(1.0, 1.4, 1.2, 1.0))
    side.paragraphs[0]._p.style = "SidebarSectionsSpacing"
    doc.add_paragraph(style="CV Photo")  # Word needs a paragraph between the last table and the section end

    # Every page ends in a yellow strip at the bottom right.
    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    foot.paragraph_format.space_after = Pt(0)
    strip = foot.add_run(" " * 22)
    strip.font.size = Pt(13)
    strip._r.get_or_add_rPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{ACCENT}"/>'))
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")  # python-docx's blank document omits this required attribute
    doc.save(Path(__file__).parent / "templates" / "bracket.docx")


if __name__ == "__main__":
    main()
