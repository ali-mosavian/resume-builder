# /// script
# requires-python = ">=3.12"
# dependencies = ["python-docx~=1.1", "pyyaml~=6.0", "typer~=0.12"]
# ///
"""Build a CV .docx from a YAML file, styled by a .docx template.

The template's last top-level table is the layout: first column the sidebar, last column the main body.
Restyle the template in Word; paragraph styles are referenced by style id below.

Optional parts, used when the template has them (see make_bracket_template.py):
placeholders CVPhoto ("{Photo <width cm> [gray]}"), CVContact and CVHighlights; styles CVRole + CVCompany
(role timeline), CVTool + CVTools (tool chips, no labels), CVSkill + CVSkillBar (level bars), CVKeyValue, CVQuote.
The CV prefix keeps them from matching a template's own styles (Word has a built-in Quote).
"""
import re
from pathlib import Path

import typer
import yaml
from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Twips

BOLD = re.compile(r"\*\*(.+?)\*\*")


def add_text(p, text: str) -> None:
    """Append text with **bold** markup as runs."""
    for i, part in enumerate(BOLD.split(text)):
        if part:
            # None, not False, outside markup: False would override bold paragraph styles.
            p.add_run(part).bold = True if i % 2 else None


def para(cell, style_id: str, text: str = ""):
    # By style id, not name: the default template has duplicate style names.
    p = cell.add_paragraph()
    p._p.style = style_id
    add_text(p, text)
    return p


def has(block, style_id: str) -> bool:
    """Whether the block's document defines the style: optional parts follow the template."""
    return block.part.styles.element.get_by_id(style_id) is not None


def section(cell, spacing_id: str, heading_id: str, title: str) -> None:
    para(cell, spacing_id)
    para(cell, heading_id, title)


def bullet_num_id(doc) -> str:
    """numId of the template's first bullet list, so bullets match the template."""
    numbering = doc.part.numbering_part.element
    # The first bullet definition, not the first instance of any bullet definition: a template's own list is
    # often registered after python-docx's built-in ones.
    first = next(a.get(qn("w:abstractNumId")) for a in numbering.findall(qn("w:abstractNum"))
                 if a.find(f"{qn('w:lvl')}/{qn('w:numFmt')}").get(qn("w:val")) == "bullet")
    return next(n.get(qn("w:numId")) for n in numbering.findall(qn("w:num"))
                if n.find(qn("w:abstractNumId")).get(qn("w:val")) == first)


def bullet(cell, text: str, num_id: str) -> None:
    p = para(cell, "ListParagraph", text)
    num_pr = OxmlElement("w:numPr")
    for tag, val in (("w:ilvl", "0"), ("w:numId", num_id)):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        num_pr.append(el)
    p._p.get_or_add_pPr().append(num_pr)


def hyperlink(p, label: str, url: str, underline: bool = True) -> None:
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), p.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = p.add_run(label)
    run.underline = underline or None
    link.append(run._r)
    p._p.append(link)


def styled(run, style_id: str):
    run._r.style = style_id
    return run


def when(ym: str | None, labels: dict) -> str:
    if ym is None:
        return labels["present"]
    year, month = str(ym).split("-")
    return f"{labels['months'][int(month) - 1]} {year}"


def span(item: dict, labels: dict) -> str:
    return f"{when(item['start'], labels)} - {when(item['end'], labels)}".upper()


def cells(tables):
    """Every cell, nested ones included; a merged cell once."""
    seen = []
    for t in tables:
        for row in t.rows:
            for c in row.cells:
                if c._tc not in seen:
                    seen.append(c._tc)
                    yield c
                    yield from cells(c.tables)


def drop(p) -> None:
    p._p.getparent().remove(p._p)


def photo(p, cv: dict) -> None:
    spec = p.text.strip("{}").split()
    for r in p.runs:
        r._r.getparent().remove(r._r)
    if cv.get("photo"):
        run = p.add_run()
        run.add_picture(str(Path(cv["_dir"]) / cv["photo"]), width=Cm(float(spec[1])))
        if "gray" in spec:
            run._r.xpath(".//a:blip")[0].append(OxmlElement("a:grayscl"))


def contact(p, cv: dict) -> None:
    for link in cv["links"]:
        q = p.insert_paragraph_before()
        q._p.style = "CVContact"
        hyperlink(q, link["label"], link["url"], underline=False)
        styled(q.add_run("  ◆"), "CVAccent")
    drop(p)


def inner_width(cell):
    """The cell's width less its left and right margins: what a nested table may fill."""
    pad = sum(int(w) for w in cell._tc.xpath("./w:tcPr/w:tcMar/*[self::w:left or self::w:right]/@w:w"))
    return cell.width - Twips(pad)


def highlights(p, cell, cv: dict) -> None:
    items = cv.get("highlights", [])
    t = cell.add_table(rows=2, cols=len(items))
    for col, c in zip(t._tbl.tblGrid.gridCol_lst, t.columns):
        col.w = int(inner_width(cell) / len(items))
        for x in c.cells:
            x.width = col.w
    for i, h in enumerate(items):
        for row, style_id, text in ((0, "CVHighlightValue", h["value"]), (1, "CVHighlightLabel", h["label"])):
            q = t.cell(row, i).paragraphs[0]
            q._p.style = style_id
            add_text(q, text)
    drop(p)


def header(doc, cv: dict) -> None:
    """Fill the template's placeholders, wherever the template puts them."""
    for p, cell in [(p, None) for p in doc.paragraphs] + [(p, c) for c in cells(doc.tables) for p in c.paragraphs]:
        match p._p.style:
            case "Name":
                p.runs[0].text = cv["name"]
            case "JobTitle":
                p.runs[0].text = cv["title"].upper()
            case "CVPhoto" if p.text.startswith("{Photo"):
                photo(p, cv)
            case "CVContact" if p.text == "{Contact}":
                contact(p, cv)
            case "CVHighlights" if p.text == "{Highlights}":
                highlights(p, cell, cv)


def bar(cell, level: int) -> None:
    """A skill level as a bar: SkillFill's shading over level/5 of the text width, SkillTrack's over the rest."""
    fill, track = (cell.part.styles.element.get_by_id(s).xpath(".//w:shd/@w:fill")[0] for s in ("CVSkillFill", "CVSkillTrack"))
    inner = inner_width(cell)
    parts = [(level, fill)] + [(5 - level, track)] * (level < 5)
    t = cell.add_table(rows=1, cols=len(parts))
    t._tbl.tblPr.get_or_add_tblLayout().set(qn("w:type"), "fixed")
    t.rows[0].height, t.rows[0].height_rule = Cm(0.12), WD_ROW_HEIGHT_RULE.EXACTLY
    for col, c, (share, color) in zip(t._tbl.tblGrid.gridCol_lst, t.rows[0].cells, parts):
        col.w = c.width = int(inner * share / 5)
        c._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{color}"/>'))
        c.paragraphs[0]._p.style = "CVSkillBar"
    cell.paragraphs[-1]._p.style = "CVSkillBar"  # the paragraph Word needs after a table


def sidebar(cell, cv: dict, lb: dict) -> None:
    section(cell, "SidebarSectionsSpacing", "Heading3", lb["details"])
    para(cell, "SidebarText", cv["location"])
    for key in ("born", "nationality"):
        para(cell, "Heading4", lb[key].upper())
        para(cell, "SidebarText", str(cv[key]))
    if not has(cell, "CVContact"):  # otherwise the header lists the links
        section(cell, "SidebarSectionsSpacing", "Heading3", lb["links"])
        for link in cv["links"]:
            hyperlink(para(cell, "SidebarText"), link["label"], link["url"])
    section(cell, "SidebarSectionsSpacing", "Heading3", lb["skills"])
    for s in cv["skills"]:
        if has(cell, "CVSkillBar"):
            para(cell, "CVSkill", s["name"])
            bar(cell, s["level"])
        else:
            para(cell, "SidebarText", f"{s['name']} · {lb['levels'][s['level']]}")
    section(cell, "SidebarSectionsSpacing", "Heading3", lb["languages"])
    for lang in cv["languages"]:
        if has(cell, "CVKeyValue"):
            styled(para(cell, "CVKeyValue", lang["name"]).add_run(f"\t{lang['level']}"), "CVMuted")
        else:
            para(cell, "SidebarText", f"{lang['name']} · {lang['level']}")
    section(cell, "SidebarSectionsSpacing", "Heading3", lb["hobbies"])
    para(cell, "SidebarText", cv["hobbies"])


def main_body(cell, cv: dict, lb: dict, num_id: str) -> None:
    cell._tc.remove(cell.paragraphs[0]._p)  # template placeholder
    para(cell, "Heading1", lb["summary"])
    for text in cv["summary"]:
        para(cell, "Normal", text)
    section(cell, "MainSectionsSpacing", "Heading1", lb["education"])
    for ed in cv["education"]:
        para(cell, "Heading2", f"{ed['degree']}, {ed['school']}, {ed['city']}")
        para(cell, "Date", span(ed, lb))
        para(cell, "Normal", ed.get("notes", ""))
    section(cell, "MainSectionsSpacing", "Heading1", lb["experience"])
    chips = has(cell, "CVTool")
    for job in cv["experience"]:
        if has(cell, "CVRole"):
            para(cell, "CVRole", job["title"])
            para(cell, "CVCompany", f"{job['company']}, {job['city']}")
        else:
            para(cell, "Heading2", f"{job['title']} {lb['at']} {job['company']}, {job['city']}")
        para(cell, "CVRoleDate" if has(cell, "CVRoleDate") else "Date", span(job, lb))
        if not chips:
            para(cell, "Normal").add_run(lb["achievements"]).italic = True
        for text in job["bullets"]:
            bullet(cell, text, num_id)
        if chips:
            p = para(cell, "CVTools")
            for tool in job["tools"]:
                # A shaded narrow no-break space draws the edge; a bar glyph made Word break chips mid-word.
                styled(p.add_run("\u202f"), "CVToolMark")
                styled(p.add_run(f"\u00a0{tool}\u00a0"), "CVTool")
                p.add_run(" \u200b")  # Word won't break at a space next to a no-break space; a zero-width space breaks
        else:
            p = para(cell, "Normal")
            p.add_run(f"{lb['tools']}: ").italic = True
            p.add_run(", ".join(job["tools"]))
    section(cell, "MainSectionsSpacing", "Heading1", lb["testimonials"])
    for t in cv["testimonials"]:
        para(cell, "Heading2", f"{t['name']} | {t['role']}")
        para(cell, "Date", when(t["date"], lb).upper())
        quote = [*t["quote"]]
        quote[0] = f"“{quote[0]}"
        quote[-1] += "”"
        for text in quote:
            para(cell, "CVQuote" if has(cell, "CVQuote") else "Normal", text)


ROOT = Path(__file__).parent


def build(cv_path: Path, out: Path | None = None, template: Path = ROOT / "templates" / "default.docx") -> None:
    out = out or ROOT / "build" / f"{cv_path.stem}_{template.stem}.docx"
    out.parent.mkdir(exist_ok=True)
    cv = yaml.safe_load(cv_path.read_text(encoding="utf-8")) | {"_dir": cv_path.parent}
    doc = Document(template)
    row = doc.tables[-1].rows[0].cells
    side, main = row[0], row[-1]
    header(doc, cv)
    sidebar(side, cv, cv["labels"])
    main_body(main, cv, cv["labels"], bullet_num_id(doc))
    doc.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    typer.run(build)
