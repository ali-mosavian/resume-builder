# /// script
# requires-python = ">=3.12"
# dependencies = ["python-docx~=1.1"]
# ///
"""Generate templates/pro.docx: name banner on top, light sidebar left, main column right.

Defines the style ids build_cv.py expects. Tweak here and regenerate, or edit the .docx in Word.
"""
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Calibri"
INK, ACCENT, MUTED, RULE, FILL = "222222", "1F3A5F", "6B7280", "C9D1DB", "F3F5F8"


def style(doc, name: str, base: str | None = None, *, size: float, color: str = INK, bold=False,
          caps=False, spacing: float = 0, before: float = 0, after: float = 2, line: float = 1.15):
    styles = doc.styles
    s = styles[name] if name in [x.name for x in styles] else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    if base:
        s.base_style = styles[base]
    f = s.font
    f.name, f.size, f.bold, f.all_caps, f.italic = FONT, Pt(size), bold, caps, False
    f.color.rgb = RGBColor.from_string(color)
    s.element.rPr.get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if spacing:
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(int(spacing * 20)))
        s.element.rPr.append(sp)
    pf = s.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line
    pf.keep_with_next = name.startswith("Heading") or name == "Date"
    return s


def bottom_rule(s, color: str, size: int = 6) -> None:
    s.element.get_or_add_pPr().append(parse_xml(
        f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{size}" w:space="2" w:color="{color}"/></w:pBdr>'))


def bullets(doc) -> None:
    """One bullet list definition; build_cv.py picks the first bullet numbering it finds."""
    numbering = doc.part.numbering_part.element
    numbering.insert(0, parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="90"><w:lvl w:ilvl="0"><w:start w:val="1"/>'
        f'<w:numFmt w:val="bullet"/><w:lvlText w:val="•"/><w:lvlJc w:val="left"/>'
        f'<w:pPr><w:ind w:left="284" w:hanging="227"/></w:pPr>'
        f'<w:rPr><w:color w:val="{ACCENT}"/></w:rPr></w:lvl></w:abstractNum>'))
    numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="90"><w:abstractNumId w:val="90"/></w:num>'))


def cell_props(cell, width_cm: float, fill: str | None = None, pad_cm: float = 0) -> None:
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    if fill:
        tc_pr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>'))
    pad = int(pad_cm * 567)
    tc_pr.append(parse_xml(
        f'<w:tcMar {nsdecls("w")}><w:top w:w="{pad}" w:type="dxa"/><w:left w:w="{pad}" w:type="dxa"/>'
        f'<w:bottom w:w="{pad}" w:type="dxa"/><w:right w:w="{pad}" w:type="dxa"/></w:tcMar>'))


def main() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.6)
    sec.top_margin = sec.bottom_margin = Cm(1.4)

    style(doc, "Normal", size=9.5, after=3)
    style(doc, "Name", "Normal", size=24, color=ACCENT, bold=True, after=0, line=1.0)
    title = style(doc, "Job Title", "Normal", size=10, color=MUTED, spacing=1.5, after=10)
    bottom_rule(title, ACCENT, 12)
    h1 = style(doc, "Heading 1", "Normal", size=10.5, color=ACCENT, bold=True, caps=True, spacing=1, before=4, after=5)
    bottom_rule(h1, RULE)
    style(doc, "Heading 2", "Normal", size=10.5, bold=True, before=6, after=0)
    style(doc, "Heading 3", "Normal", size=9, color=ACCENT, bold=True, caps=True, spacing=1, after=3)
    style(doc, "Heading 4", "Normal", size=7.5, color=MUTED, caps=True, spacing=0.8, before=3, after=0)
    style(doc, "Date", "Normal", size=8, color=MUTED, spacing=0.5, after=3)
    style(doc, "Sidebar Text", "Normal", size=9, after=1)
    style(doc, "List Paragraph", "Normal", size=9.5, after=2)
    style(doc, "Sidebar Sections Spacing", "Normal", size=4, after=6)
    style(doc, "Main Sections Spacing", "Normal", size=4, after=4)
    bullets(doc)

    doc.add_paragraph(style="Name").add_run("{Name}")
    doc.add_paragraph(style="Job Title").add_run("{Job Title}")

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (5.0, 0.5, 12.3)  # sums to the 17.8 cm text width
    # Word and LibreOffice size columns from the grid, not from cell widths alone.
    for col, w in zip(table._tbl.tblGrid.gridCol_lst, widths):
        col.w = Cm(w)
    table._tbl.tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
    side, gap, body = table.rows[0].cells
    cell_props(side, widths[0], FILL, 0.35)
    cell_props(gap, widths[1])
    cell_props(body, widths[2])
    side.paragraphs[0]._p.style = "SidebarSectionsSpacing"
    doc.save(Path(__file__).parent / "templates" / "pro.docx")


if __name__ == "__main__":
    main()
