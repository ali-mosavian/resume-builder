import re

import pytest
from docx import Document
from docx.shared import Cm

from support import CV
from resume import data
from support import docx
from resume.word import cells
from support import TEMPLATES
from resume.docx_templates import names
from resume.docx_templates import bold_font

TEMPLATE_FILES = sorted(p.stem for p in TEMPLATES.glob("*.docx"))


@pytest.mark.parametrize("template", TEMPLATE_FILES)
def test_template_fills_in(template: str) -> None:
    doc = docx(template)
    main = doc.tables[-1].rows[0].cells[-1]
    # Plain text once got an explicit bold=False, which un-bolded every heading.
    assert not any(r.bold is False for p in main.paragraphs for r in p.runs), "explicit bold=False found"
    # **markup** becomes a bold run without the asterisks.
    first = next(p for p in main.paragraphs if "shelf allocation engine" in p.text)
    assert "**" not in first.text
    assert any(
        (r.bold or r._r.style == "CVStrong") and r.text == "shelf allocation engine" for r in first.runs
    )
    assert "SEP 2025 - PRESENT" in [p.text for p in main.paragraphs]
    everywhere = [*doc.paragraphs, *(p for c in cells(doc.tables) for p in c.paragraphs)]
    assert any(p.text == "Alex Morgan" for p in everywhere), "name placeholder not filled"
    assert not [p.text for p in everywhere if p.text.startswith("{")], "placeholders left"
    # Optional parts once keyed on common style names (SkillBar, Quote) that other templates also have:
    # the resume.io template crashed and pro restyled its testimonials. Neither opts in.
    quote = next(p for p in main.paragraphs if p.text.startswith("“It is rare"))
    assert (quote._p.style == "CVQuote") == (template not in ("default", "pro"))


@pytest.mark.parametrize(("template", "color"), [("pro", "1F3A5F"), ("bracket", "F6B92B")])
def test_bullets_use_the_template_list(template: str, color: str) -> None:
    """Bullets once came from python-docx's built-in list, not the template's own: black instead of the accent colour."""
    doc = docx(template)
    num_id = doc.element.body.xpath(".//w:numPr/w:numId/@w:val")[0]
    numbering = doc.part.numbering_part.element
    abstract = numbering.xpath(f"./w:num[@w:numId='{num_id}']/w:abstractNumId/@w:val")[0]
    assert numbering.xpath(
        f"./w:abstractNum[@w:abstractNumId='{abstract}']/w:lvl[@w:ilvl='0']/w:rPr/w:color/@w:val"
    ) == [color]


def test_no_dashes() -> None:
    """Dates were printed "JUN 2024 <em dash> NU": the CV uses "-", never em or en dashes."""
    found = re.search(".{20}[–—].{20}", docx("bracket", "sv").element.xml)
    assert not found, found[0]


def test_tool_chips_can_wrap_between_chips() -> None:
    """Word glues a no-break space to the spaces around it, so tool chips (padded with no-break spaces) had no break
    point between them and Word broke chips mid-word ("SQLAlchem|y"). Each separator carries a zero-width space."""
    chips = next(p for p in docx("bracket").tables[-1].rows[0].cells[-1].paragraphs if p._p.style == "CVTools")
    assert chips.text.count("\u200b") == chips.text.count(" ")


def test_bracket_word_layout() -> None:
    """In Word (LibreOffice hid both): the Word 2010 layout mode aligned the layout table's text, not its edge, with the
    margin, pushing the sidebar off the page; and two heading styles sharing one list lost the sidebar's diamonds."""
    bracket = Document(str(TEMPLATES / "bracket.docx"))
    mode = bracket.settings.element.xpath("./w:compat/w:compatSetting[@w:name='compatibilityMode']/@w:val")
    assert mode == ["15"]
    lists = [
        s.element.xpath("./w:pPr/w:numPr/w:numId/@w:val")[0]
        for s in bracket.styles
        if s.element.xpath("./w:pPr/w:numPr/w:numId")
    ]
    assert len(lists) == len(set(lists)), f"styles share a list: {lists}"


def test_bracket_font() -> None:
    """bracket.docx once set its headings in Calibri: Word's heading styles name a theme font, which beats the named one.
    Its text was 222222 while bracket.css inks 262626. And Montserrat travels inside the file, for readers without it."""
    bracket = Document(str(TEMPLATES / "bracket.docx"))
    assert not [
        s.style_id for s in bracket.styles if s.element.xpath("./w:rPr/w:rFonts[@w:ascii='Montserrat'][@w:asciiTheme]")
    ]
    assert bracket.styles["Normal"].element.xpath("./w:rPr/w:color/@w:val") == ["262626"]
    embedded = [p.partname for p in docx("bracket", "sv").part.package.iter_parts() if p.partname.endswith(".odttf")]
    assert len(embedded) == 6  # Montserrat's four faces, and Montserrat Bold's two (see test_bold_is_a_real_bold_face)


def test_pro_grid_widths() -> None:
    """template_pro once had only cell widths, no grid widths: renderers split the page evenly,
    squeezing the main column to a third and doubling the page count."""
    grid = Document(str(TEMPLATES / "pro.docx")).tables[0]._tbl.tblGrid.gridCol_lst
    assert (grid[2].w or 0) > Cm(12)


@pytest.mark.parametrize("template", names())
def test_bold_is_a_real_bold_face(template: str) -> None:
    """Word for Mac ignores embedded bold faces and fakes bold from the regular one: thin "bold" in every theme whose
    font Office doesn't ship. Bold styles use the font's bold family instead, and **bold** text uses CVStrong."""
    styles = Document(str(TEMPLATES / f"{template}.docx")).styles
    faked = [
        s.style_id
        for s in styles
        if s.element.xpath("./w:rPr/w:b[not(@w:val) or @w:val='1' or @w:val='true']")
        and (font := s.font.name)
        and bold_font(font)
    ]
    assert not faked, f"styles faking bold: {faked}"
    assert "CV Strong" in [s.name for s in styles]


def test_split_is_one_column() -> None:
    """split.docx's layout is a single cell: the body first, then the sidebar's sections, skills as chips."""
    doc = docx("split")
    (body,) = doc.tables[-1].rows[0].cells
    texts = [p.text for p in body.paragraphs]
    labels = data.load(CV, "en")["labels"]
    assert texts.index(labels["summary"]) < texts.index(labels["testimonials"]) < texts.index(labels["skills"])
    skills = next(p for p in body.paragraphs if p._p.style == "CVSkillChips")
    assert "Python" in skills.text


def test_split_photo_is_round() -> None:
    assert 'prst="ellipse"' in docx("split").element.xml


def test_split_band_is_not_dimmed() -> None:
    """Word dims header content while the body is edited: split's slate band, drawn in the first-page header, showed
    grey in Word. It is anchored in the body instead, where it keeps its colour."""
    doc = Document(str(TEMPLATES / "split.docx"))
    assert not doc.sections[0].first_page_header._element.xpath(".//wp:anchor")
    assert doc.element.body.xpath(".//wp:anchor//a:srgbClr[@val='0F172A']")
