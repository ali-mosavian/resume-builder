# /// script
# requires-python = ">=3.12"
# dependencies = ["python-docx~=1.1", "pyyaml~=6.0", "typer~=0.12"]
# ///
"""Self-check: builds cv/cv_en.yaml with each template and checks what the reader sees. Run: uv run test_build_cv.py"""
import re
import shutil
import subprocess
import tempfile

import yaml
from pathlib import Path

from docx import Document
from docx.shared import Cm

from build_cv import build, cells

here = Path(__file__).parent

# A bold edit once quoted only the end of a Swedish bullet, printing a stray ' after Nordea's text.
for path in (here / "cv").glob("*.yaml"):
    for j in yaml.safe_load(path.read_text(encoding="utf-8"))["experience"]:
        for b in j["bullets"]:
            assert b.endswith("."), f"{path.name}, {j['company']}: bullet ends {b[-12:]!r}"
for template in sorted((here / "templates").glob("*.docx")):
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "cv.docx"
        build(here / "cv" / "cv_en.yaml", out, template)
        doc = Document(out)
        main = doc.tables[-1].rows[0].cells[-1]
        texts = [p.text for p in main.paragraphs]

        # Plain text once got an explicit bold=False, which un-bolded every heading.
        assert not any(r.bold is False for p in main.paragraphs for r in p.runs), "explicit bold=False found"
        # **markup** becomes a bold run without the asterisks.
        first = next(p for p in main.paragraphs if "store capacity optimization engine" in p.text)
        assert "**" not in first.text and any(r.bold and r.text == "store capacity optimization engine" for r in first.runs)
        assert "SEP 2025 — PRESENT" in texts
        everywhere = [*doc.paragraphs, *(p for c in cells(doc.tables) for p in c.paragraphs)]
        assert any(p.text == "Ali Mosavian" for p in everywhere), f"{template.name}: name placeholder not filled"
        left = [p.text for p in everywhere if p.text.startswith("{")]
        assert not left, f"{template.name}: placeholders left: {left}"
        # Optional parts once keyed on common style names (SkillBar, Quote) that other templates also have:
        # the resume.io template crashed and pro restyled its testimonials. Only bracket opts in.
        quote = next(p for p in main.paragraphs if p.text.startswith("“It is rare"))
        assert (quote._p.style == "CVQuote") == (template.stem == "bracket"), f"{template.name}: quote styled {quote._p.style}"

# Bullets once came from python-docx's built-in list, not the template's own: black instead of the accent colour.
for template, color in (("pro", "1F3A5F"), ("bracket", "F6B92B")):
    with tempfile.TemporaryDirectory() as tmp:
        build(here / "cv" / "cv_en.yaml", Path(tmp) / "cv.docx", here / "templates" / f"{template}.docx")
        doc = Document(Path(tmp) / "cv.docx")
        num_id = doc.element.body.xpath(".//w:numPr/w:numId/@w:val")[0]
        numbering = doc.part.numbering_part.element
        abstract = numbering.xpath(f"./w:num[@w:numId='{num_id}']/w:abstractNumId/@w:val")[0]
        used = numbering.xpath(f"./w:abstractNum[@w:abstractNumId='{abstract}']/w:lvl[@w:ilvl='0']/w:rPr/w:color/@w:val")
        assert used == [color], f"{template}: bullets coloured {used}, template says {color}"

# bracket.docx once set its headings in Calibri: Word's heading styles name a theme font, which beats the named one.
# And Montserrat travels inside the file, for readers who don't have it.
bracket = Document(here / "templates" / "bracket.docx")
themed = [s.style_id for s in bracket.styles if s.element.xpath("./w:rPr/w:rFonts[@w:ascii='Montserrat'][@w:asciiTheme]")]
assert not themed, f"bracket.docx styles shadowed by a theme font: {themed}"
with tempfile.TemporaryDirectory() as tmp:
    build(here / "cv" / "cv_sv.yaml", Path(tmp) / "cv.docx", here / "templates" / "bracket.docx")
    embedded = [p.partname for p in Document(Path(tmp) / "cv.docx").part.package.iter_parts() if p.partname.endswith(".odttf")]
    assert len(embedded) == 4, f"bracket CV embeds {len(embedded)} font files"

# template_pro once had only cell widths, no grid widths: renderers split the page evenly,
# squeezing the main column to a third and doubling the page count.
grid = Document(here / "templates" / "pro.docx").tables[0]._tbl.tblGrid.gridCol_lst
assert grid[2].w > Cm(12), "template_pro main column is not sized in the table grid"
with tempfile.TemporaryDirectory() as tmp:
    from build_html import build as build_html

    for theme in sorted(p.stem for p in (here / "themes").glob("*.css")):
        out = Path(tmp) / f"cv_{theme}.html"
        build_html(here / "cv" / "cv_en.yaml", out, theme)
        html = out.read_text(encoding="utf-8")
        assert "<strong>store capacity optimization engine</strong>" in html and "**" not in html
        assert "reading &amp; learning" in html, "YAML text is not HTML-escaped"
        # Brief roles are not contiguous (Watty sits between them): the caption must appear once.
        assert html.count("data-earlier=") == 1, "earlier-roles caption repeated or missing"
    # One page, both languages: the second starts hidden; the theme script and the language switch are inlined.
    out = Path(tmp) / "cv_both.html"
    build_html(here / "cv" / "cv_en.yaml", out, "aurora-v3", also=[here / "cv" / "cv_sv.yaml"])
    html = out.read_text(encoding="utf-8")
    assert '<div class="page" lang="en"' in html and '<div class="page" lang="sv"' in html
    assert re.search(r'<div class="page" lang="sv" data-ui="[^"]*" hidden>', html), "second language not hidden"
    assert not re.search(r'<div class="page" lang="en" data-ui="[^"]*" hidden>', html), "first language hidden"
    assert html.count("<script>") == 2 and 'data-aka="GCP"' in html
    # Printing during the headline count-up once put partial numbers ("2+", "7M") in the PDF.
    chrome = Path("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")
    if chrome.exists() and shutil.which("pdftotext"):
        page, pdf = Path(tmp) / "cv_print.html", Path(tmp) / "cv.pdf"
        build_html(here / "cv" / "cv_en.yaml", page, "aurora-v3")
        subprocess.run([chrome, "--headless=new", "--disable-gpu", "--virtual-time-budget=300", "--no-pdf-header-footer",
                        f"--print-to-pdf={pdf}", page.as_uri()], capture_output=True, check=True)
        text = subprocess.run(["pdftotext", "-l", "1", pdf, "-"], capture_output=True, text=True, check=True).stdout.split()
        assert {"20+", "17", "42+"} <= set(text), f"printed headline numbers are partial: {text[:30]}"
    else:
        print("skipped print check: needs Google Chrome and pdftotext")

    # slate.js paginates into A4 sheets that clip overflow; chained testimonials once overflowed one sheet
    # and the last four vanished from the PDF. Every role and testimonial must reach the paper, on A4.
    if chrome.exists() and shutil.which("pdftotext"):
        cv = yaml.safe_load((here / "cv" / "cv_en.yaml").read_text(encoding="utf-8"))
        pager = (here / "themes" / "slate.js").resolve()
        for theme in sorted(p.stem for p in (here / "themes").glob("*.js") if p.resolve() == pager):
            page, pdf = Path(tmp) / f"cv_{theme}.html", Path(tmp) / f"cv_{theme}.pdf"
            build_html(here / "cv" / "cv_en.yaml", page, theme)
            subprocess.run([chrome, "--headless=new", "--disable-gpu", "--window-size=1280,1000", "--virtual-time-budget=4000",
                            "--no-pdf-header-footer", f"--print-to-pdf={pdf}", page.as_uri()], capture_output=True, check=True)
            # Raw (content-stream) order: the default reading order once wove sidebar lines into a wrapped role heading
            # and reported Nordea lost. Raw drops spaces, so compare without whitespace.
            text = re.sub(r"\s", "", subprocess.run(["pdftotext", "-raw", pdf, "-"], capture_output=True, text=True, check=True).stdout)
            names = [*(j["company"] for j in cv["experience"]), *(t["name"] for t in cv["testimonials"])]
            missing = [n for n in names if re.sub(r"\s", "", n) not in text]
            assert not missing, f"{theme} print lost: {missing}"
            info = subprocess.run(["pdfinfo", pdf], capture_output=True, text=True, check=True).stdout
            assert "(A4)" in info, f"{theme} print is not A4"

    # bracket once dropped the sidebar column on sheets after its content ran out: every sheet keeps it.
    if chrome.exists():
        probe = ("<script>addEventListener('load',()=>setTimeout(()=>{document.title=[...document.querySelectorAll('.sheet aside')]"
                 ".map(a=>a.getBoundingClientRect().width|0).join(',')},1500))</script>")
        page = Path(tmp) / "cv_dmc_side.html"
        build_html(here / "cv" / "cv_en.yaml", page, "bracket")
        page.write_text(page.read_text(encoding="utf-8").replace("</body>", probe + "</body>"), encoding="utf-8")
        dom = subprocess.run([chrome, "--headless=new", "--disable-gpu", "--window-size=1280,1000", "--virtual-time-budget=4000",
                              "--dump-dom", page.as_uri()], capture_output=True, text=True, check=True).stdout
        widths = [int(w) for w in re.search("<title>(.*?)</title>", dom)[1].split(",")]
        assert len(widths) > 1 and min(widths) > 100, f"bracket sidebar widths per sheet: {widths}"

    # Language switch on a paginated theme: the language opened by #sv is cut into sheets, and so is
    # the one switched to; the download links follow the YAML.
    if chrome.exists():
        probe = ("<script>addEventListener('load',()=>setTimeout(()=>{const n=()=>[...document.querySelectorAll('.page')]"
                 ".map(p=>p.lang+(p.hidden?'-':'+')+p.querySelectorAll('.sheet').length).join(' ');const a=n();"
                 "const x=!!document.querySelector('.page[lang=sv] .sheet:first-child [data-from=experience]');"
                 "document.querySelector('.page:not([hidden]) .lang-switch [data-lang=en]').click();"
                 "setTimeout(()=>{const y=!!document.querySelector('.page[lang=en] .sheet:first-child [data-from=experience]');"
                 "document.title=a+'|'+n()+'|'+x+'|'+y},1000)},1500))</script>")
        page = Path(tmp) / "cv_switch.html"
        build_html(here / "cv" / "cv_en.yaml", page, "bracket", also=[here / "cv" / "cv_sv.yaml"])
        html = page.read_text(encoding="utf-8")
        assert 'href="ali_mosavian-resume-sv.docx" download' in html and "onclick=\"print()\"" not in html
        page.write_text(html.replace("</body>", probe + "</body>"), encoding="utf-8")
        dom = subprocess.run([chrome, "--headless=new", "--disable-gpu", "--window-size=1280,1000", "--virtual-time-budget=6000",
                              "--dump-dom", page.as_uri() + "#sv"], capture_output=True, text=True, check=True).stdout
        before, after, sv_first, en_first = re.search("<title>(.*?)</title>", dom)[1].split("|")
        # Experience once jumped to sheet 2, leaving a third of sheet 1 empty: in Swedish because the role's bottom
        # padding was counted, in English because a long first bullet could not break across sheets.
        assert sv_first == en_first == "true", f"Experience not on sheet 1: sv {sv_first}, en {en_first}"
        assert before.startswith("en-0 sv+") and not before.endswith("+0"), f"opened #sv: {before}"
        assert after.startswith("en+") and not after.startswith("en+0"), f"switched to en: {after}"

    # "Show details" on earlier roles once turned their tool tags into a plain block (display:block beat flex):
    # expanded, they must lay out like any other role's tags.
    if chrome.exists():
        probe = ("<script>addEventListener('load',()=>{const d=e=>getComputedStyle(e).display;"
                 "document.querySelector('.page:not([hidden]) .details-toggle').click();"
                 "document.title=d(document.querySelector('.item:not(.brief) .tools'))+'|'+d(document.querySelector('.brief .tools'))})</script>")
        shared = (here / "themes" / "aurora-v3.js").resolve()
        for theme in sorted(p.stem for p in (here / "themes").glob("*.js") if p.resolve() == shared):
            page = Path(tmp) / f"cv_expand_{theme}.html"
            build_html(here / "cv" / "cv_en.yaml", page, theme)
            page.write_text(page.read_text(encoding="utf-8").replace("</body>", probe + "</body>"), encoding="utf-8")
            dom = subprocess.run([chrome, "--headless=new", "--disable-gpu", "--virtual-time-budget=1000", "--dump-dom", page.as_uri()],
                                 capture_output=True, text=True, check=True).stdout
            normal, brief = re.search("<title>(.*?)</title>", dom)[1].split("|")
            assert normal == brief, f"{theme}: expanded earlier-role tools are {brief}, other roles' are {normal}"
print("ok")
