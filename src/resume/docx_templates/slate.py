"""templates/slate.docx, the Word twin of themes/slate.css.

Slate sidebar on every page (a shape in the page headers) with photo, name and title on top; white main column.
"""

from pathlib import Path
from functools import partial

from docx import Document
from docx.shared import Cm
from docx.shared import Pt
from docx.document import Document as Doc
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

from resume.docx_templates import Mark
from resume.docx_templates import cell
from resume.docx_templates import rect
from resume.docx_templates import style
from resume.docx_templates import table
from resume.docx_templates import border
from resume.docx_templates import finish
from resume.docx_templates import char_style
from resume.docx_templates import bullet_lists

FONT = "Lato"
INK, MUTED, SIDE, SIDE_MUTED, TRACK, WHITE, ROLE = "1D1F23", "8A8F99", "4C546A", "A3A9B8", "6B7287", "FFFFFF", "D9DCE4"
LINE_GAP = 1.2  # Lato's line height per em: Word's "multiple" spacing scales it, CSS line-height doesn't

SIDE_W = 68  # sidebar width in mm, as themes/slate.css lays out an A4 sheet
TOP = 12  # main column's top padding; page 1's header sits 2 mm lower


def styles(doc: Doc) -> None:
    font = partial(style, doc, base="Normal", font=FONT, color=INK, line=1.5, line_gap=LINE_GAP)

    bullet_lists(doc, [Mark("•", INK, 360, 240)])

    font("Normal", base=None, size=9.5, before=6, after=6)
    centered = WD_ALIGN_PARAGRAPH.CENTER
    font("CV Photo", size=2, after=0, line=1.0).paragraph_format.alignment = centered
    font(
        "Name", size=18.6, color=WHITE, bold=True, spacing=0.19, after=0, line=1.5
    ).paragraph_format.alignment = centered
    rule = font("CV Rule", size=1, color=SIDE_MUTED, before=10.5, after=12, line=1.0)
    rule.paragraph_format.left_indent = rule.paragraph_format.right_indent = Cm(2.26)  # an 18px divider
    border(rule, "bottom", SIDE_MUTED, 6, 0)
    job = font("Job Title", size=7.9, color=ROLE, caps=True, spacing=1.58, after=0, line=2.0)
    job.paragraph_format.alignment = centered

    font("Heading 1", size=15.6, bold=True, before=0, after=0)
    font("Heading 2", size=11.4, bold=True, before=15, after=0)
    font("Date", size=7.9, color=MUTED, caps=True, spacing=1.1, before=1.5, after=1.5)
    font("List Paragraph", size=9.5, after=1.5).paragraph_format.left_indent = Cm(0.635)
    font("CV Quote", size=9.5, after=6)

    font("Heading 3", size=12.96, color=WHITE, bold=True, before=0, after=7.5)
    font("Heading 4", size=7.44, color=SIDE_MUTED, bold=True, caps=True, spacing=1.19, before=9, after=1.5)
    font("Sidebar Text", size=10.3, color=WHITE, after=4)
    font("CV Skill", size=10.3, color=WHITE, after=3.75)
    font("CV Skill Bar", size=1, after=9, line=1.0)
    kv = font("CV Key Value", size=10.3, color=WHITE, after=4.5)
    kv.paragraph_format.tab_stops.add_tab_stop(Cm(5.0), alignment=WD_TAB_ALIGNMENT.RIGHT)  # the sidebar text edge
    font("Sidebar Sections Spacing", size=1, after=24.3, line=1.0)
    font("Main Sections Spacing", size=1, after=18, line=1.0)
    char_style(doc, "CV Muted", font=FONT, color=SIDE_MUTED, size=9.4)
    char_style(doc, "CV Skill Fill", font=FONT, fill=WHITE)
    char_style(doc, "CV Skill Track", font=FONT, fill=TRACK)
    for s in ("Heading 1", "Heading 2", "Heading 3", "Heading 4", "Date"):
        doc.styles[s].paragraph_format.keep_with_next = True


def layout(doc: Doc) -> None:
    """Body: sidebar column (its slate is a shape, see decorate) with the header on top, and main column."""
    side, main = table(doc, [SIDE_W / 10, 21 - SIDE_W / 10]).rows[0].cells
    cell(side, SIDE_W / 10, pad=(0.2, 0.9, 0, 0.9))
    cell(main, 21 - SIDE_W / 10, pad=(0, 1.1, 0, 1.1))
    photo = side.paragraphs[0]
    photo._p.style = "CVPhoto"
    photo.add_run("{Photo 2.0}")
    photo.paragraph_format.space_after = Pt(10.5)
    side.add_paragraph(style="Name").add_run("{Name}")
    side.add_paragraph(style="CV Rule")
    side.add_paragraph(style="Job Title").add_run("{Job Title}")
    main.paragraphs[0]._p.style = "CVPhoto"
    doc.add_paragraph(style="CV Photo")  # Word needs a paragraph between the last table and the section end


def decorate(doc: Doc) -> None:
    """The slate sidebar behind the text, repeated by the page headers."""
    sec = doc.sections[0]
    for i, header in enumerate((sec.first_page_header, sec.header)):
        p = header.paragraphs[0]
        p._p.style = "CVPhoto"
        rect(p, (0, 0, SIDE_W, 297, SIDE), 1001 + i)


def make(out: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(0)
    sec.top_margin, sec.bottom_margin = Cm(TOP / 10), Cm(1.2)
    sec.header_distance, sec.footer_distance = Cm(0), Cm(0)
    sec.different_first_page_header_footer = True
    styles(doc)
    layout(doc)
    decorate(doc)
    finish(doc, out, FONT)
