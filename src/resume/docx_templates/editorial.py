"""templates/editorial.docx, the Word twin of themes/editorial.css.

Warm paper as a page shape, a ruled masthead, "No. 01" section heads, role titles over italic company lines,
tool chips, quotes. The web's drop cap and the collapsed earlier roles have no Word counterpart.
"""

from pathlib import Path
from functools import partial

from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn
from docx.document import Document as Doc
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

from resume.ooxml import xml
from resume.docx_templates import Mark
from resume.docx_templates import cell
from resume.docx_templates import rect
from resume.docx_templates import style
from resume.docx_templates import table
from resume.docx_templates import border
from resume.docx_templates import finish
from resume.docx_templates import numbered
from resume.docx_templates import char_style
from resume.docx_templates import bullet_lists

SERIF, SANS = "Newsreader", "Inter"
PAPER, PAPER_2, INK, MUTED, RULE, RULE_2, ACCENT = "F6F1E8", "EFE8DC", "1A1714", "625B52", "D9D1C3", "C9BFAE", "A8321A"
SERIF_GAP, SANS_GAP = 1.0, 1.21  # hhea line height per em: Word's "multiple" scales it, CSS line-height doesn't

# Page geometry in mm, as themes/editorial.css lays out an A4 sheet.
MARGIN_X, SIDE_W, GAP, MAIN_W = 14, 48, 9, 125
HEADING_NUM_ID = "95"


def styles(doc: Doc) -> None:
    serif = partial(style, doc, base="Normal", font=SERIF, color=INK, line_gap=SERIF_GAP)
    sans = partial(style, doc, base="Normal", font=SANS, color=INK, line_gap=SANS_GAP)

    bullet_lists(doc, [Mark("-", ACCENT, 220, 220, 9, SERIF)])  # word.py's bullets take the first
    numbered_heads(doc)

    serif("Normal", base=None, size=9.5, line=1.45, after=7)
    serif("CV Gap", size=1, line=1.0, after=0)
    serif("CV Header Gap", size=1, line=1.0, after=17)
    sans("CV Kicker", size=7, color=MUTED, caps=True, spacing=1.5, before=0.5, after=11.5, line=1.0)
    serif("Name", size=36, spacing=-0.9, before=12, after=0, line=1.0)
    job = serif("Job Title", size=12, color=ACCENT, caps=True, spacing=1.0, before=6, after=0, line=1.25)
    job.font.italic = True
    sans("CV Contact", size=7.5, before=0, after=0, line=1.5).paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    serif("CV Highlights", size=1, line=1.0, after=0)
    serif("CV Highlight Value", size=18, spacing=-0.5, before=7, after=3, line=1.0)
    sans("CV Highlight Label", size=6, color=MUTED, caps=True, spacing=0.9, after=6, line=1.42)

    h1 = serif("Heading 1", size=14, before=0, after=8, line=1.2)
    border(h1, "top", INK, 12, 6)
    numbered(h1, HEADING_NUM_ID)
    serif("Heading 2", size=11.5, before=0, after=0, line=1.3)
    serif("CV Role", base="Heading 2", size=11.5, before=2, after=0, line=1.3)
    serif("CV Company", size=9, color=MUTED, after=0, line=1.4).font.italic = True
    sans("Date", size=6.5, color=MUTED, caps=True, spacing=0.9, before=2, after=5, line=1.6)
    sans("CV Role Date", base="Date", size=6.5, color=ACCENT, caps=True, spacing=0.9, before=2, after=5, line=1.6)
    h3 = sans("Heading 3", size=7, caps=True, spacing=1.0, before=0, after=9.5, line=1.2)
    border(h3, "bottom", INK, 6, 6)
    sans("Heading 4", size=6, color=MUTED, caps=True, spacing=0.8, before=5, after=0, line=1.4)
    serif("Sidebar Text", size=8, line=1.5, after=0)
    kv = serif("CV Key Value", size=8, line=1.5, before=1, after=0)
    kv.paragraph_format.tab_stops.add_tab_stop(Cm(SIDE_W / 10), alignment=WD_TAB_ALIGNMENT.RIGHT)
    border(kv, "between bottom", RULE, 4, 3)
    serif("CV Skill", size=8, line=1.5, after=2)
    serif("CV Skill Bar", size=1, line=1.0, after=6)
    serif("List Paragraph", size=9, line=1.42, after=2)
    tools = sans("CV Tools", size=6.5, line=1.8, before=5, after=7)
    border(tools, "bottom", RULE, 4, 8)
    quote = serif("CV Quote", size=9.5, line=1.45, after=10)
    quote.font.italic = True
    serif("Sidebar Sections Spacing", size=1, line=1.0, after=9)
    serif("Main Sections Spacing", size=1, line=1.0, after=8)
    char_style(doc, "CV Accent", font=SANS, color=ACCENT, size=5)
    muted = char_style(doc, "CV Muted", font=SANS, color=MUTED, size=6)
    muted.font.all_caps = True
    char_style(doc, "CV Tool", font=SANS, color=MUTED, fill=PAPER_2, size=6.5)
    char_style(doc, "CV Tool Mark", font=SANS, size=6.5)  # unshaded: only spaces the chips apart
    char_style(doc, "CV Skill Fill", font=SERIF, fill=INK)
    char_style(doc, "CV Skill Track", font=SERIF, fill=RULE)
    for s in ("CV Role", "CV Company", "Date", "CV Role Date", "Heading 1", "Heading 2", "Main Sections Spacing"):
        doc.styles[s].paragraph_format.keep_with_next = True
    doc.styles["List Paragraph"].paragraph_format.keep_with_next = False

    # Nested tables (highlights) get no side padding from the default table style.
    for mar in doc.styles.element.xpath(
        "./w:style[@w:styleId='TableNormal']//w:tblCellMar/*[self::w:left or self::w:right]"
    ):
        mar.set(qn("w:w"), "0")


def numbered_heads(doc: Doc) -> None:
    """ "No. 01" in the margin column of Heading 1: the number, then a tab to the title."""
    numbering = doc.part.numbering_part.element
    numbering.insert(
        0,
        xml(
            f'<w:abstractNum w:abstractNumId="{HEADING_NUM_ID}"><w:lvl w:ilvl="0"><w:start w:val="1"/>'
            '<w:numFmt w:val="decimalZero"/><w:suff w:val="tab"/><w:lvlText w:val="No. %1"/><w:lvlJc w:val="left"/>'
            '<w:pPr><w:tabs><w:tab w:val="num" w:pos="794"/></w:tabs><w:ind w:left="794" w:hanging="794"/></w:pPr>'
            f'<w:rPr><w:rFonts w:ascii="{SANS}" w:hAnsi="{SANS}"/><w:caps/><w:color w:val="{ACCENT}"/>'
            '<w:spacing w:val="20"/><w:sz w:val="14"/></w:rPr></w:lvl></w:abstractNum>'
        ),
    )
    numbering.append(xml(f'<w:num w:numId="{HEADING_NUM_ID}"><w:abstractNumId w:val="{HEADING_NUM_ID}"/></w:num>'))


def header(doc: Doc) -> None:
    """Masthead: kicker, name and role left, links right, headline numbers ruled below."""
    head = table(doc, [(182 - 45) / 10, 4.5], rows=2)
    left, right = head.rows[0].cells
    rule = {"top": (INK, 16)}
    cell(left, 13.7, pad=(0, 0, 0.49, 0), borders=rule)
    cell(right, 4.5, pad=(0.21, 0, 0, 0), borders=rule)
    left.paragraphs[0]._p.style = "CVKicker"
    left.paragraphs[0].add_run("Curriculum Vitæ")
    left.add_paragraph(style="Name").add_run("{Name}")
    left.add_paragraph(style="JobTitle").add_run("{Job Title}")
    right.paragraphs[0]._p.style = "CVContact"
    right.paragraphs[0].add_run("{Contact}")
    stats = head.rows[1].cells[0].merge(head.rows[1].cells[1])
    cell(stats, 18.2, borders={"top": (INK, 8), "bottom": (RULE_2, 6)})
    stats.paragraphs[0]._p.style = "CVHighlights"
    stats.paragraphs[0].add_run("{Highlights}")


def layout(doc: Doc) -> None:
    doc.add_paragraph(style="CV Header Gap")  # keeps Word from merging the two tables
    side, main = table(doc, [(SIDE_W + GAP) / 10, MAIN_W / 10]).rows[0].cells
    cell(side, (SIDE_W + GAP) / 10, pad=(0, GAP / 10, 0, 0))
    cell(main, MAIN_W / 10)
    side.paragraphs[0]._p.style = "CVGap"  # word.py adds a section spacer of its own
    doc.add_paragraph(style="CV Gap")  # Word needs a paragraph between the last table and the section end


def paper(doc: Doc) -> None:
    """The page colour: a shape behind the text, repeated by both page headers."""
    sec = doc.sections[0]
    for i, head in enumerate((sec.first_page_header, sec.header)):
        head.paragraphs[0]._p.style = "CVGap"
        rect(head.paragraphs[0], (0, 0, 210, 297, PAPER), 1001 + i)


def make(out: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(MARGIN_X / 10)
    sec.top_margin, sec.bottom_margin = Cm(1.3), Cm(1.4)
    sec.header_distance, sec.footer_distance = Cm(0), Cm(0)
    sec.different_first_page_header_footer = True
    styles(doc)
    header(doc)
    layout(doc)
    paper(doc)
    finish(doc, out, SERIF, SANS)
