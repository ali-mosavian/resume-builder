"""templates/kinetic.docx, the Word twin of themes/kinetic.css.

Beyond the style ids every template has, it opts into word.py's optional parts: Contact and Highlights placeholders
(no photo), skill bars, a role headline with the company beneath, tool chips and quote bars.
"""

from pathlib import Path
from functools import partial

from docx import Document
from docx.shared import Cm
from docx.document import Document as Doc
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ROW_HEIGHT_RULE

from resume.docx_templates import Mark
from resume.docx_templates import cell
from resume.docx_templates import rect
from resume.docx_templates import style
from resume.docx_templates import table
from resume.docx_templates import border
from resume.docx_templates import finish
from resume.docx_templates import char_style
from resume.docx_templates import bullet_lists

FONT, DISPLAY = "Inter", "Bricolage Grotesque"
PAPER, INK, COBALT, AMBER = "F3EFE6", "141412", "1C2CD9", "F2C230"
MUTED, SKY, TINT, TRACK, WHITE = "5F5C55", "C9CEFF", "E6E2D8", "CFCABE", "FFFFFF"
SANS_GAP, DISPLAY_GAP = (
    1.21,
    1.2,
)  # font line height per em: Word's "multiple" spacing scales it, CSS line-height doesn't

# Page geometry in mm, as themes/kinetic.css lays out an A4 sheet.
ROW1, ROW2, STRIP, SIDE_W = 42, 30, 8, 62  # hero: name and contact, headline numbers (strip at its foot); sidebar width
HERO = ROW1 + ROW2


def styles(doc: Doc) -> None:
    sans = partial(style, doc, base="Normal", font=FONT, color=INK, line=1.55, line_gap=SANS_GAP)
    display = partial(style, doc, base="Normal", font=DISPLAY, color=INK, line_gap=DISPLAY_GAP)
    (bullet,) = bullet_lists(doc, [Mark("•", COBALT, 283, 227, 10, "Segoe UI Symbol")])
    assert bullet

    sans("Normal", base=None, size=9, after=4)
    display("Name", size=60, bold=True, caps=True, color=WHITE, spacing=-1.5, line=0.9, after=0)
    sans("Job Title", size=6.8, bold=True, caps=True, color=AMBER, spacing=0.95, line=1.0, after=7)
    sans("CV Photo", size=2, after=0, line=1.0)
    contact = sans("CV Contact", size=8, bold=True, color=WHITE, line=1.0, after=0)
    contact.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    border(contact, "top between bottom", WHITE, 8, 4)
    sans("CV Highlights", size=2, after=0)
    display("CV Highlight Value", size=24, bold=True, color=AMBER, spacing=-0.5, line=1.1, after=3)
    sans("CV Highlight Label", size=7, caps=True, color=SKY, spacing=0.7, line=1.3, after=0)

    heading = sans("Heading 1", size=8, bold=True, caps=True, spacing=1.28, line=1.2, before=0, after=10)
    border(heading, "top", INK, 8, 6)
    side_heading = sans("Heading 3", size=8, bold=True, caps=True, spacing=1.28, line=1.2, before=0, after=10)
    border(side_heading, "top", INK, 8, 6)
    sans("Heading 2", size=9.5, bold=True, before=0, after=0)
    display("CV Role", base="Heading 2", size=12, bold=True, spacing=-0.24, line=1.15, before=0, after=0)
    sans("CV Company", size=8.2, bold=True, color=COBALT, line=1.4, after=3)
    sans("Date", size=7, caps=True, color=MUTED, spacing=0.84, line=1.4, after=6)
    sans("CV Role Date", base="Date", size=7, caps=True, color=MUTED, spacing=0.84, line=1.4, after=9)
    sans("Heading 4", size=6.6, bold=True, caps=True, color=MUTED, spacing=0.79, line=1.4, before=7, after=1)
    sans("Sidebar Text", size=8.5, after=1.4)
    kv = sans("CV Key Value", size=8.5, after=4.5)
    kv.paragraph_format.tab_stops.add_tab_stop(Cm(4.3), alignment=WD_TAB_ALIGNMENT.RIGHT)  # at the sidebar text edge
    sans("CV Skill", size=8.5, after=2)
    sans("CV Skill Bar", size=1, after=7, line=1.0)
    sans("List Paragraph", size=9, after=2)
    sans("CV Tools", size=7.2, after=15, line=1.8)
    quote = sans("CV Quote", size=9, color="3A3A3A", after=4)
    quote.font.italic = True
    quote.paragraph_format.left_indent = Cm(0.48)
    border(quote, "left", COBALT, 18, 11)
    sans("Sidebar Sections Spacing", size=4, after=17, line=1.0)
    sans("Main Sections Spacing", size=4, after=23, line=1.0)
    char_style(doc, "CV Accent", font=FONT, color=AMBER, size=5.5)
    char_style(doc, "CV Muted", font=FONT, color=MUTED, size=7.8)
    char_style(doc, "CV Tool", font=FONT, fill=TINT, size=7.2)
    char_style(doc, "CV Tool Mark", font=FONT, fill=COBALT, size=7.2)  # a chip's cobalt left edge
    char_style(doc, "CV Skill Fill", font=FONT, fill=COBALT)
    char_style(doc, "CV Skill Track", font=FONT, fill=TRACK)
    for s in ("CV Contact", "CV Role", "CV Company", "Date", "CV Role Date", "Heading 1", "Heading 3"):
        doc.styles[s].paragraph_format.keep_with_next = True


def header(doc: Doc) -> None:
    """Hero rows of fixed height, so page 1's sidebar shape starts right under the amber strip."""
    head = table(doc, [14.5, 6.5], rows=2)
    for row, mm in zip(head.rows, (ROW1, ROW2), strict=True):
        row.height, row.height_rule = Cm(mm / 10), WD_ROW_HEIGHT_RULE.EXACTLY
    name_c, contact_c = head.rows[0].cells
    cell(name_c, 14.5, pad=(1.2, 0, 0, 1.2))
    cell(contact_c, 6.5, pad=(1.2, 1.4, 0, 0))
    name_c.paragraphs[0]._p.style = "JobTitle"
    name_c.paragraphs[0].add_run("{Job Title}")
    name_c.add_paragraph(style="Name").add_run("{Name}")
    contact_c.paragraphs[0]._p.style = "CVContact"
    contact_c.paragraphs[0].add_run("{Contact}")
    stats = head.rows[1].cells[0].merge(head.rows[1].cells[1])
    cell(stats, 21, pad=(0.5, 0.8, 0, 1.0))
    stats.paragraphs[0]._p.style = "CVHighlights"
    stats.paragraphs[0].add_run("{Highlights}")


def layout(doc: Doc) -> None:
    """Body: sidebar column (its tint is a shape, see decorate) and main column."""
    doc.add_paragraph(style="CV Photo")  # keeps Word from merging the two tables
    side, main = table(doc, [SIDE_W / 10, 21 - SIDE_W / 10]).rows[0].cells
    cell(side, SIDE_W / 10, pad=(0, 0.7, 0, 1.2))
    cell(main, 21 - SIDE_W / 10, pad=(1.2, 1.4, 0, 1.0))
    side.paragraphs[0]._p.style = "CVPhoto"  # word.py adds a section spacer of its own
    doc.add_paragraph(style="CV Photo")  # Word needs a paragraph between the last table and the section end


def decorate(doc: Doc) -> None:
    """Page decoration as shapes behind the text, repeated by the page headers: the paper, then on page 1 the cobalt
    hero with its amber strip and rule, and the sidebar tint from under it (later pages: from the top, with the
    body starting 14 mm down as on the web's later sheets)."""
    sec = doc.sections[0]
    first = sec.first_page_header.paragraphs[0]
    first._p.style = "CVPhoto"
    for i, box in enumerate(
        [
            (0, 0, 210, 297, PAPER),
            (0, 0, 210, HERO, COBALT),
            (0, HERO - STRIP, 210, STRIP, AMBER),
            (12, ROW1, 186, 0.3, WHITE),
            (0, HERO, SIDE_W, 297 - HERO, TINT),
        ]
    ):
        rect(first, box, 1001 + i)
    later = sec.header.paragraphs[0]
    later._p.style = "CVPhoto"
    later.paragraph_format.space_before = Cm(0.1)
    for i, box in enumerate([(0, 0, 210, 297, PAPER), (0, 0, SIDE_W, 297, TINT)]):
        rect(later, box, 1101 + i)


def make(out: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(0)
    sec.top_margin, sec.bottom_margin = Cm(0), Cm(1.2)
    sec.header_distance, sec.footer_distance = Cm(0), Cm(0)
    sec.different_first_page_header_footer = True
    styles(doc)
    header(doc)
    layout(doc)
    decorate(doc)
    finish(doc, out, FONT, DISPLAY)
