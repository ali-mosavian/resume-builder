"""templates/split.docx, the Word twin of themes/split.css on paper: one column on A4, not the web's two halves.

A slate band on page 1 holds the photo, name, links and headline numbers; the body follows in one column, and the
layout table's single cell makes word.py close the document with the sidebar's sections. Opts into word.py's
Photo, Contact and Highlights placeholders, tool chips, skill chips and quote bars.
"""

from pathlib import Path
from functools import partial

from docx import Document
from docx.shared import Cm
from docx.document import Document as Doc
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

from resume.docx_templates import Mark
from resume.docx_templates import cell
from resume.docx_templates import rect
from resume.docx_templates import style
from resume.docx_templates import table
from resume.docx_templates import border
from resume.docx_templates import finish
from resume.docx_templates import char_style
from resume.docx_templates import bullet_lists

FONT = "Inter"
SLATE, HEAD, INK, BODY, FAINT = "0F172A", "0F172A", "1E293B", "334155", "94A3B8"
TEAL, TEAL_LIGHT, CHIP, LINE, ON_SLATE, ON_SLATE_2 = "0F766E", "5EEAD4", "E6F2F1", "CBD5E1", "F1F5F9", "CBD5E1"
LINE_GAP = 1.21  # Inter's line height per em: Word's "multiple" spacing scales it, CSS line-height doesn't

# Page geometry in mm, as themes/split.css prints an A4 sheet.
HEAD1, HEAD2 = 41, 24  # band rows: photo, name and links; then the headline numbers
BAND = HEAD1 + HEAD2 + 3  # the slate band, a little past the numbers
MARGIN = 1.6  # body margin left and right, cm


def styles(doc: Doc) -> None:
    font = partial(style, doc, base="Normal", font=FONT, color=INK, line=1.5, line_gap=LINE_GAP)
    bullet_lists(doc, [Mark("•", TEAL, 255, 255)])  # word.py's bullets take the first

    font("Normal", base=None, size=9.5, after=6)
    font("Name", size=24, color=ON_SLATE, bold=True, spacing=-0.4, after=0, line=1.1)
    font("Job Title", size=8, color=ON_SLATE_2, caps=True, spacing=1.0, before=3, after=0)
    font("CV Photo", size=2, after=0, line=LINE_GAP)  # single spacing: below it LibreOffice clips the photo
    contact = font("CV Contact", size=8, color="E2E8F0", after=2, line=1.2)
    contact.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font("CV Highlights", size=2, after=0)
    font("CV Highlight Value", size=15, color=TEAL_LIGHT, bold=True, after=0, line=1.1)
    font("CV Highlight Label", size=7, color=FAINT, after=0, line=1.3)

    # Section heads: small, tracked, teal, over a hairline. The sidebar's (Heading 3) look the same: one column.
    for name in ("Heading 1", "Heading 3"):
        border(font(name, size=8, color=TEAL, bold=True, caps=True, spacing=1.1, after=7), "bottom", LINE, 4, 3)
    font("Heading 2", size=10, color=HEAD, bold=True, before=6, after=0)
    font("Heading 4", size=7, color=FAINT, caps=True, spacing=0.6, before=4, after=0)
    font("Date", size=7, color=TEAL, bold=True, caps=True, spacing=0.4, before=1, after=4)
    font("Sidebar Text", size=9, after=1)
    kv = font("CV Key Value", size=9, after=1)
    kv.paragraph_format.tab_stops.add_tab_stop(Cm(4.5), alignment=WD_TAB_ALIGNMENT.LEFT)
    font("List Paragraph", size=9, after=2)
    font("CV Tools", size=7, before=4, after=8, line=1.9)
    font("CV Skill Chips", size=7, before=0, after=2, line=1.9)
    quote = font("CV Quote", size=9, color=BODY, after=6)
    quote.font.italic = True
    border(quote, "left", TEAL_LIGHT, 12, 8)
    font("Sidebar Sections Spacing", size=4, after=10)
    font("Main Sections Spacing", size=4, after=10)
    char_style(doc, "CV Accent", font=FONT, color=TEAL_LIGHT, size=6)
    char_style(doc, "CV Muted", font=FONT, color=FAINT)
    char_style(doc, "CV Tool", font=FONT, color=TEAL, fill=CHIP, size=7)
    char_style(doc, "CV Tool Mark", font=FONT, fill=CHIP, size=7)
    for s in ("Heading 1", "Heading 2", "Heading 3", "Date"):
        doc.styles[s].paragraph_format.keep_with_next = True


def header(doc: Doc) -> None:
    """Band rows of fixed height, so the body starts right under the slate."""
    head = table(doc, [4.6, 10.4, 6.0], rows=2)
    for row, mm in zip(head.rows, (HEAD1, HEAD2), strict=True):
        row.height, row.height_rule = Cm(mm / 10), WD_ROW_HEIGHT_RULE.EXACTLY
    photo_c, name_c, links_c = head.rows[0].cells
    cell(photo_c, 4.6, pad=(1.1, 0, 0, MARGIN))
    cell(name_c, 10.4, pad=(1.1, 0.4, 0, 0.3), valign=WD_CELL_VERTICAL_ALIGNMENT.CENTER)
    cell(links_c, 6.0, pad=(1.1, MARGIN, 0, 0.4), valign=WD_CELL_VERTICAL_ALIGNMENT.CENTER)
    photo_c.paragraphs[0]._p.style = "CVPhoto"
    photo_c.paragraphs[0].add_run("{Photo 2.8}")
    name_c.paragraphs[0]._p.style = "Name"
    name_c.paragraphs[0].add_run("{Name}")
    name_c.add_paragraph(style="Job Title").add_run("{Job Title}")
    links_c.paragraphs[0]._p.style = "CVContact"
    links_c.paragraphs[0].add_run("{Contact}")
    stats = head.rows[1].cells[0].merge(head.rows[1].cells[2])
    cell(stats, 21, pad=(0.4, MARGIN, 0, MARGIN), borders={"top": ("334155", 4)})
    stats.paragraphs[0]._p.style = "CVHighlights"
    stats.paragraphs[0].add_run("{Highlights}")


def layout(doc: Doc) -> None:
    """Body: one cell, so word.py writes the main body and then the sidebar's sections into it."""
    doc.add_paragraph(style="CV Photo")  # keeps Word from merging the two tables
    (body,) = table(doc, [21]).rows[0].cells
    cell(body, 21, pad=(0.9, MARGIN, 0, MARGIN))
    # word.py drops the cell's first paragraph; this second one stays first, so the cell doesn't open with a
    # keep-with-next heading (LibreOffice then moves the whole table to page 2).
    body.add_paragraph(style="CV Photo")
    doc.add_paragraph(style="CV Photo")  # Word needs a paragraph between the last table and the section end


def decorate(doc: Doc) -> None:
    """The band as a shape behind page 1's header rows; later pages start 12 mm down, as the PDF's do."""
    sec = doc.sections[0]
    first = sec.first_page_header.paragraphs[0]
    first._p.style = "CVPhoto"
    rect(first, (0, 0, 210, BAND, SLATE), 1001)
    later = sec.header.paragraphs[0]
    later._p.style = "CVPhoto"
    later.paragraph_format.space_before = Cm(1.2)


def make(out: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(0)
    sec.top_margin, sec.bottom_margin = Cm(0), Cm(1.2)
    sec.header_distance, sec.footer_distance = Cm(0), Cm(0)
    sec.different_first_page_header_footer = True
    styles(doc)
    header(doc)
    layout(doc)
    decorate(doc)
    finish(doc, out, FONT)
