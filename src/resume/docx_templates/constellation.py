"""templates/constellation.docx, the Word twin of themes/constellation.css.

Night sky over the first page's header (shapes in the first-page header), paper below. The web page draws the
career as a star chart on a canvas; Word gets a fixed, abstract one made of shapes. Opts into word.py's optional
parts like bracket does: Contact and Highlights placeholders, skill bars, a role timeline, tool chips, quote bars.
"""

import random
from pathlib import Path
from functools import partial

from docx import Document
from docx.shared import Cm
from docx.shared import Pt
from docx.document import Document as Doc
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.enum.table import WD_ROW_HEIGHT_RULE

from resume.ooxml import xml
from resume.docx_templates import Box
from resume.docx_templates import WPS
from resume.docx_templates import Mark
from resume.docx_templates import cell
from resume.docx_templates import rect
from resume.docx_templates import style
from resume.docx_templates import table
from resume.docx_templates import border
from resume.docx_templates import finish
from resume.docx_templates import numbered
from resume.docx_templates import char_style
from resume.docx_templates import bullet_lists

SERIF, SANS, MONO = "Instrument Serif", "IBM Plex Sans", "IBM Plex Mono"
INK, INK2, MUTED, LINE = "141827", "2C3346", "5B6378", "E1E4EE"
GOLD, BLUE, BLUE_SOFT, PAPER, TRACK = "E8BD63", "3B47C9", "E9EBFB", "F6F7FB", "DFE2EF"
SKY, STAR, DIM, RULE = "0B1030", "CFD8FF", "AAB3CF", "2C3252"
LINE_GAP = 1.3  # all three families: hhea line height per em

# Page geometry in mm, as themes/constellation.css lays out an A4 sheet.
HEAD1, HEAD2, SIDE_W = 85, 25, 56  # header rows (name block; headline numbers) and sidebar width
HEAD = HEAD1 + HEAD2


def styles(doc: Doc) -> None:
    font = partial(style, doc, base="Normal", font=SANS, color=INK2, line=1.5, line_gap=LINE_GAP)
    mono = partial(font, font=MONO)
    serif = partial(font, font=SERIF)

    _, ring = bullet_lists(  # word.py's bullets take the first
        doc, [Mark("•", GOLD, 567, 227, 10, "Segoe UI Symbol"), Mark("○", BLUE, 340, 340, 12, "Segoe UI Symbol")]
    )

    font("Normal", base=None, size=9, after=8.5)
    serif("Name", size=38, color="FFFFFF", after=17, line=1.0)
    mono("Job Title", size=7, color=GOLD, caps=True, spacing=1.0, after=8.5)
    font("CV Photo", size=2, after=0, line=1.0)
    contact = font("CV Contact", size=8, color="E5E9F7", before=0, after=0)
    contact.paragraph_format.right_indent = Cm(2.4)  # the web's links are 56 mm wide
    border(contact, "top between bottom", RULE, 6, 4)
    font("CV Highlights", size=2, after=0)
    value = font("CV Highlight Value", size=20, font=SERIF, color="FFFFFF", after=1.4, line=1.1)
    label = font("CV Highlight Label", size=7.2, color=DIM, after=0, line=1.4)
    for s_ in (value, label):
        s_.paragraph_format.right_indent = Cm(0.4)  # the web's gap between columns
    border(value, "top", RULE, 6, 7)

    border(serif("Heading 1", size=19, color=INK, before=0, after=11.3, line=1.2), "bottom", LINE, 6, 4)
    font("Heading 2", size=9.6, bold=True, color=INK, before=0, after=0)
    numbered(font("CV Role", base="Heading 2", size=9.6, bold=True, color=INK, before=0, after=0), ring)
    serif("CV Company", size=11.5, color=MUTED, after=0, line=1.3).paragraph_format.left_indent = Cm(0.6)
    mono("Date", size=7, color=MUTED, caps=True, spacing=0.42, before=1.7, after=5.7, line=1.5)
    role_date = mono("CV Role Date", base="Date", size=7, color=MUTED, caps=True, spacing=0.42, before=1.7, after=5.7)
    role_date.paragraph_format.left_indent = Cm(0.6)  # under the role title, past the ring
    mono("Heading 3", size=7, color=MUTED, caps=True, spacing=1.12, before=0, after=8.5, line=1.4)
    mono("Heading 4", size=6.6, color=MUTED, caps=True, spacing=0.8, before=6.8, after=0.85, line=1.4)
    font("Sidebar Text", size=8.5, after=1.4)
    kv = font("CV Key Value", size=8.5, after=4.5)
    kv.paragraph_format.tab_stops.add_tab_stop(Cm(3.4), alignment=WD_TAB_ALIGNMENT.RIGHT)  # at the sidebar text edge
    font("CV Skill", size=8.5, color=INK, after=2.8)
    font("CV Skill Bar", size=1, after=5.8, line=1.0)
    font("List Paragraph", size=9, after=2.3).paragraph_format.left_indent = Cm(1.0)
    font("CV Tools", size=6.8, after=11.3, line=1.8).paragraph_format.left_indent = Cm(0.6)
    quote = font("CV Quote", size=9, after=10)
    quote.font.italic = True
    quote.paragraph_format.left_indent = Cm(0.4)
    border(quote, "left", GOLD, 18, 8)
    font("Sidebar Sections Spacing", size=4, after=14, line=1.0)
    font("Main Sections Spacing", size=4, after=16, line=1.0)
    char_style(doc, "CV Accent", font=SANS, color=GOLD)
    char_style(doc, "CV Muted", font=SANS, color=MUTED, size=7.3)
    char_style(doc, "CV Tool", font=MONO, fill=BLUE_SOFT, size=6.8, color=INK2)
    char_style(doc, "CV Tool Mark", font=MONO, fill=GOLD, size=6.8)  # a chip's gold left edge
    char_style(doc, "CV Skill Fill", font=SANS, fill=BLUE)
    char_style(doc, "CV Skill Track", font=SANS, fill=TRACK)
    for s in ("CV Contact", "CV Role", "CV Company", "Date", "CV Role Date", "Heading 1", "Heading 3"):
        doc.styles[s].paragraph_format.keep_with_next = True


def header(doc: Doc) -> None:
    """Header rows of fixed height, so page 1's sidebar shape starts right under the headline numbers."""
    head = table(doc, [9.4, 11.6], rows=2)
    for row, mm in zip(head.rows, (HEAD1, HEAD2), strict=True):
        row.height, row.height_rule = Cm(mm / 10), WD_ROW_HEIGHT_RULE.EXACTLY
    text, _ = head.rows[0].cells
    cell(text, 9.4, pad=(1.2, 0, 0, 1.4))
    cell(head.rows[0].cells[1], 11.6)
    text.paragraphs[0]._p.style = "JobTitle"
    text.paragraphs[0].add_run("{Job Title}")
    text.add_paragraph(style="Name").add_run("{Name}")
    text.add_paragraph(style="CV Contact").add_run("{Contact}")
    stats = head.rows[1].cells[0].merge(head.rows[1].cells[1])
    cell(stats, 21, pad=(0.4, 1.4, 0, 1.4))
    stats.paragraphs[0]._p.style = "CVHighlights"
    stats.paragraphs[0].add_run("{Highlights}")


def layout(doc: Doc) -> None:
    """Body: sidebar column (its paper tint is a shape, see decorate) and main column."""
    # Keeps Word from merging the two tables. Its space is the gap under the header: a cell's top padding would
    # repeat on every page.
    doc.add_paragraph(style="CV Photo").paragraph_format.space_after = Pt(36)
    side, main = table(doc, [SIDE_W / 10, 21 - SIDE_W / 10]).rows[0].cells
    cell(side, SIDE_W / 10, pad=(0, 0.8, 0, 1.4))
    cell(main, 21 - SIDE_W / 10, pad=(0, 1.4, 0, 1.0))
    side.paragraphs[0]._p.style = "CVPhoto"  # word.py adds a section spacer of its own
    doc.add_paragraph(style="CV Photo")  # Word needs a paragraph between the last table and the section end


def shape(
    p: Paragraph,
    kind: str,
    box: tuple[float, float, float, float],
    color: str,
    shape_id: int,
    *,
    alpha: int = 100,
    line_mm: float = 0,
    flip: bool = False,
) -> None:
    """An ellipse (filled) or a line (stroked, from the box's top-left to bottom-right, or bottom-left with flip)."""
    x, y, w, h = box

    def emu(mm: float) -> int:
        return int(mm * 36000)

    fill = f'<a:solidFill><a:srgbClr val="{color}"><a:alpha val="{alpha * 1000}"/></a:srgbClr></a:solidFill>'
    if kind == "line":
        paint = f'<a:noFill/><a:ln w="{emu(line_mm)}">{fill}</a:ln>'
    else:
        paint = f"{fill}<a:ln><a:noFill/></a:ln>"
    p._p.append(
        xml(
            f'<w:r xmlns:wps="{WPS}"><w:drawing>'
            f'<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="{shape_id}" behindDoc="1" '
            f'locked="1" layoutInCell="0" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
            f'<wp:positionH relativeFrom="page"><wp:posOffset>{emu(x)}</wp:posOffset></wp:positionH>'
            f'<wp:positionV relativeFrom="page"><wp:posOffset>{emu(y)}</wp:posOffset></wp:positionV>'
            f'<wp:extent cx="{emu(w)}" cy="{emu(h)}"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
            f'<wp:docPr id="{shape_id}" name="Shape {shape_id}"/><wp:cNvGraphicFramePr/>'
            f'<a:graphic><a:graphicData uri="{WPS}"><wps:wsp><wps:cNvSpPr/><wps:spPr>'
            f'<a:xfrm{' flipV="1"' if flip else ""}><a:off x="0" y="0"/><a:ext cx="{emu(w)}" cy="{emu(h)}"/></a:xfrm>'
            f'<a:prstGeom prst="{"line" if kind == "line" else "ellipse"}"><a:avLst/></a:prstGeom>{paint}</wps:spPr><wps:bodyPr/></wps:wsp>'
            f"</a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>",
            "w",
            "wp",
            "a",
        )
    )


def sky(p: Paragraph) -> None:
    """A fixed star chart in the header's right half: roles on an axis, stars above them, linked."""
    rng = random.Random(7)
    sid = 1100
    x0, x1, axis, top = 96.0, 194.0, 76.0, 17.0
    roles = [x0 + (x1 - x0) * (1 - (1 - i / 9) ** 1.6) for i in range(10)]
    stars = [(rng.uniform(x0 + 4, x1 - 6), rng.uniform(top, axis - 9), rng.uniform(0.9, 2.0)) for _ in range(22)]
    shape(p, "line", (x0 - 2, axis, x1 - x0 + 4, 0), DIM, sid, alpha=45, line_mm=0.2)
    for r in roles:
        sid += 1
        shape(p, "line", (r, axis - 1.6, 0, 1.6), DIM, sid, alpha=50, line_mm=0.2)
    for sx, sy, _ in stars:
        for r in rng.sample(roles, rng.choice((1, 2, 2, 3))):
            sid += 1
            shape(
                p,
                "line",
                (min(sx, r), sy, abs(r - sx) or 0.01, axis - sy),
                "8FA0E6",
                sid,
                alpha=22,
                line_mm=0.15,
                flip=r < sx,
            )
    for r in roles:
        sid += 1
        shape(p, "ellipse", (r - 0.7, axis - 0.7, 1.4, 1.4), STAR, sid)
    for sx, sy, d in stars:
        sid += 1
        shape(p, "ellipse", (sx - d * 1.7, sy - d * 1.7, d * 3.4, d * 3.4), STAR, sid, alpha=18)
        sid += 1
        shape(p, "ellipse", (sx - d / 2, sy - d / 2, d, d), STAR, sid)
    for _ in range(50):  # dust over the whole sky
        sid += 1
        d = rng.uniform(0.3, 0.7)
        shape(
            p,
            "ellipse",
            (rng.uniform(2, 208), rng.uniform(2, HEAD - 2), d, d),
            "FFFFFF",
            sid,
            alpha=rng.randint(25, 60),
        )


def decorate(doc: Doc) -> None:
    """Page decoration as shapes behind the text, repeated by the page headers. Page 1: the night sky over the header
    and the sidebar tint below it. Later pages: the tint from the top, and the body starts 14 mm down, as the web's
    later sheets do."""
    sec = doc.sections[0]
    first = sec.first_page_header.paragraphs[0]
    first._p.style = "CVPhoto"
    boxes: list[Box] = [(0, 0, 210, HEAD, SKY), (0, HEAD, SIDE_W, 297 - HEAD, PAPER)]
    for i, box in enumerate(boxes):
        rect(first, box, 1001 + i)
    sky(first)
    later = sec.header.paragraphs[0]
    later._p.style = "CVPhoto"
    later.paragraph_format.space_before = Cm(1.3)
    rect(later, (0, 0, SIDE_W, 297, PAPER), 1901)


def make(out: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(0)
    sec.top_margin, sec.bottom_margin = Cm(0), Cm(1.2)
    sec.header_distance, sec.footer_distance = Cm(0), Cm(0)
    sec.different_first_page_header_footer = True
    styles(doc)
    header(doc)
    layout(doc)
    decorate(doc)
    finish(doc, out, SERIF, SANS, MONO)
