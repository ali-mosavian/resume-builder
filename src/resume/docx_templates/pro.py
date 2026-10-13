"""templates/pro.docx, the Word twin of themes/pro.css: name banner on top, light sidebar left, main column right.

Beyond the style ids every template has, it opts into word.py's optional parts: CVKeyValue (languages) and
bordered tool chips. Testimonials stay "Normal".
"""

from pathlib import Path
from functools import partial

from docx import Document
from docx.shared import Cm
from docx.document import Document as Doc
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.enum.table import WD_ROW_HEIGHT_RULE

from resume.ooxml import xml
from resume.docx_templates import Mark
from resume.docx_templates import cell
from resume.docx_templates import rect
from resume.docx_templates import style
from resume.docx_templates import table
from resume.docx_templates import border
from resume.docx_templates import finish
from resume.docx_templates import char_style
from resume.docx_templates import bullet_lists

FONT = "Inter"
INK, ACCENT, MUTED, RULE, SIDE, CHIP = "1F2328", "1F3A5F", "6B7280", "D5DBE3", "F3F5F8", "374151"
LINE_GAP = 1.21  # Inter's line height per em: Word's "multiple" spacing scales it, CSS line-height doesn't

# Sheet geometry in mm, as themes/pro.css lays out an A4 sheet.
HEAD, SIDE_W = 34, 62  # name banner height (with its rule); sidebar width
ANCHOR = 0.9  # the paragraph between the banner and the body tables (Word would merge them otherwise)


def styles(doc: Doc) -> None:
    font = partial(style, doc, base="Normal", font=FONT, color=INK, line=1.5, line_gap=LINE_GAP)
    bullet_lists(doc, [Mark("•", ACCENT, 255, 255)])

    font("Normal", base=None, size=9.5, after=5.7)
    font("Name", size=24, color=ACCENT, bold=True, spacing=-0.24, after=0, line=1.1)
    font("Job Title", size=8, color=MUTED, caps=True, spacing=1.12, before=4.3, after=0)
    font("Anchor", size=2, after=0, line=1.0)
    border(font("Heading 1", size=8.5, color=ACCENT, bold=True, caps=True, spacing=1.19, after=7), "bottom", RULE, 6, 4)
    font("Heading 2", size=10, bold=True, before=8.5, after=0)
    font("Heading 3", size=7.5, color=ACCENT, bold=True, caps=True, spacing=1.05, after=5.7)
    font("Heading 4", size=6.5, color=MUTED, caps=True, spacing=0.65, before=5.7, after=0)
    font("Date", size=7, color=MUTED, caps=True, spacing=0.56, before=1.4, after=5.7)
    font("Sidebar Text", size=8.5, after=0)
    kv = font("CV Key Value", size=8.5, after=0)
    kv.paragraph_format.tab_stops.add_tab_stop(Cm((SIDE_W - 13) / 10), alignment=WD_TAB_ALIGNMENT.RIGHT)
    font("List Paragraph", size=9.5, after=2.3)
    font("CV Tools", size=7, after=0, line=1.9).paragraph_format.space_before = Cm(0.2)
    font("Sidebar Sections Spacing", size=4, after=8)
    font("Main Sections Spacing", size=4, after=5)
    char_style(doc, "CV Muted", font=FONT, color=MUTED)
    chip = char_style(doc, "CV Tool", font=FONT, color=CHIP, size=7)
    chip.element.get_or_add_rPr().append(xml(f'<w:bdr w:val="single" w:sz="6" w:space="0" w:color="{RULE}"/>'))
    char_style(doc, "CV Tool Mark", font=FONT, size=2)
    for s in ("Heading 1", "Heading 2", "Heading 3", "Date"):
        doc.styles[s].paragraph_format.keep_with_next = True


def layout(doc: Doc) -> None:
    # Three grid columns merged into one cell: tests expect the first table's third column to be the main one.
    banner = table(doc, [SIDE_W / 10, 0.4, 14.4])
    row = banner.rows[0]
    row.height, row.height_rule = Cm(HEAD / 10), WD_ROW_HEIGHT_RULE.EXACTLY
    head = row.cells[0].merge(row.cells[2])
    cell(head, 21, pad=(0.9, 1.4, 0, 1.4), borders={"bottom": (ACCENT, 18)})
    head.paragraphs[0]._p.style = "Name"
    head.paragraphs[0].add_run("{Name}")
    head.add_paragraph(style="Job Title").add_run("{Job Title}")
    doc.add_paragraph(style="Anchor")
    side, main = table(doc, [SIDE_W / 10, 21 - SIDE_W / 10]).rows[0].cells
    cell(side, SIDE_W / 10, pad=(0.7, 0.6, 0, 0.7))
    cell(main, 21 - SIDE_W / 10, pad=(0.7, 1.2, 0, 0.8))
    side.paragraphs[0]._p.style = "SidebarSectionsSpacing"
    doc.add_paragraph(style="Anchor")  # Word needs a paragraph between the last table and the section end


def decorate(doc: Doc) -> None:
    """The sidebar's tint as a shape behind the text: from under the banner on page 1, from the top after."""
    sec = doc.sections[0]
    first = sec.first_page_header.paragraphs[0]
    first._p.style = "Anchor"
    y = HEAD + ANCHOR
    rect(first, (0, y, SIDE_W, 297 - y, SIDE), 1001)
    later = sec.header.paragraphs[0]
    later._p.style = "Anchor"
    later.paragraph_format.space_before = Cm(1.1)  # the body starts 12 mm down, as the web's later sheets do
    rect(later, (0, 0, SIDE_W, 297, SIDE), 1101)


def make(out: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(0)
    sec.top_margin, sec.bottom_margin = Cm(0), Cm(1.0)
    sec.header_distance, sec.footer_distance = Cm(0), Cm(0)
    sec.different_first_page_header_footer = True
    styles(doc)
    layout(doc)
    decorate(doc)
    finish(doc, out, FONT)
