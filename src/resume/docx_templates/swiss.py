"""templates/swiss.docx, the Word twin of themes/swiss.css.

Word has no data column or lead paragraph, so the CSS stacks roles (title, company, dates) as Word writes them.
Optional parts of word.py used: Highlights placeholder, skill bars, role titles, tool chips and quote bars; no
Contact (the sidebar lists the links) and no Photo.
"""

from pathlib import Path
from functools import partial

from docx import Document
from docx.shared import Cm
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.enum.text import WD_UNDERLINE
from docx.document import Document as Doc
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.styles.style import ParagraphStyle
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

from resume.ooxml import xml
from resume.docx_templates import Mark
from resume.docx_templates import cell
from resume.docx_templates import rect
from resume.docx_templates import style
from resume.docx_templates import table
from resume.docx_templates import border
from resume.docx_templates import finish
from resume.docx_templates import numbered
from resume.docx_templates import char_style
from resume.docx_templates import bullet_lists

FONT = "Inter Tight"
INK, MUTED, LINE, HAIR, RED, BODY = "111111", "5C5C5C", "D9D9D9", "ECECEC", "E2231A", "262626"
LINE_GAP = 1.21  # Inter Tight's line height per em: Word's "multiple" spacing scales it, CSS line-height doesn't

# Page geometry in mm, as themes/swiss.css lays out an A4 sheet.
M, SIDE_W, GAP = 13, 44, 6  # page margin, sidebar width, gutter
SHIFT = 4.8  # Word sets the name and headline numbers this much lower than Chrome; the CSS is moved down to match
CONTACT, NAME, STATS = 9, 38 + SHIFT, 24 - 3.5  # header rows: top rule, name block, headline numbers
TOP = 12
STATS_Y = TOP + CONTACT + NAME  # the red bar; the rule sits 5 mm below
BODY_Y = STATS_Y + STATS
ROLE_X = 139
WIDTH = 210 - 2 * M
NAME_W = 75


def styles(doc: Doc) -> None:
    font = partial(style, doc, base="Normal", font=FONT, color=INK, line=1.4, line_gap=LINE_GAP)

    # One list per style: Word links a list to a single style and drops the marks of any other.
    bullet_lists(doc, [Mark("-", INK, 255, 255, 9)])
    # Section numbers: a list of its own, "01" in red ahead of each Heading 1.
    doc.part.numbering_part.element.insert(
        0,
        xml(
            '<w:abstractNum w:abstractNumId="70"><w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimalZero"/>'
            '<w:lvlText w:val="%1"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="340" w:hanging="340"/></w:pPr>'
            f'<w:rPr><w:color w:val="{RED}"/></w:rPr></w:lvl></w:abstractNum>'
        ),
    )
    doc.part.numbering_part.element.append(xml('<w:num w:numId="70"><w:abstractNumId w:val="70"/></w:num>'))

    font("Normal", base=None, size=9, after=4)
    name = font("Name", size=44, bold=True, spacing=-1.98, after=0, line=0.86)
    name.paragraph_format.line_spacing = Pt(37.84)  # exactly: Word shrinks a multiple below 1 from the top
    font("Job Title", size=9.5, bold=True, caps=True, spacing=0.57, after=0, line=1.3)
    font("CV Photo", size=2, after=0, line=1.0)
    font("CV Highlights", size=2, after=0)
    # Negative indent: the nested table's cell margin would push the text off its column edge.
    value = font("CV Highlight Value", size=18, bold=True, spacing=-0.72, after=4, line=1.0)
    label = font("CV Highlight Label", size=6.8, color=MUTED, caps=True, spacing=0.82, after=0, line=1.3)
    for s in (value, label):
        s.paragraph_format.left_indent = Cm(-0.19)

    def rule(s: ParagraphStyle) -> ParagraphStyle:
        border(s, "top", INK, 6, 5)
        return s

    label_font = partial(font, size=6.8, caps=True, spacing=0.82, line=1.3)
    numbered(rule(label_font("Heading 1", before=0, after=8)), "70")
    rule(label_font("Heading 3", before=0, after=8))
    font("Heading 2", size=9.5, bold=True, after=0, line=1.2)
    font("CV Role", base="Heading 2", size=9.5, bold=True, after=0, line=1.2)
    font("CV Company", size=8.2, color=MUTED, after=0, line=1.2)
    date = font("Date", size=7, spacing=0.28, caps=True, before=2, after=5, line=1.3)
    date.font.underline = WD_UNDERLINE.THICK
    date.element.get_or_add_rPr().xpath("./w:u")[0].set(qn("w:color"), RED)
    font("Heading 4", size=6.4, color=MUTED, bold=True, caps=True, spacing=0.77, before=6, after=0, line=1.3)
    font("Sidebar Text", size=8, after=0)
    kv = font("CV Key Value", size=8, before=2, after=2)
    kv.paragraph_format.tab_stops.add_tab_stop(Cm(SIDE_W / 10), alignment=WD_TAB_ALIGNMENT.RIGHT)
    border(kv, "between bottom", HAIR, 2, 2)
    font("CV Skill", size=8, after=3.4)
    font("CV Skill Bar", size=1, after=2.2, line=1.0)
    font("List Paragraph", size=9, color=BODY, after=2).paragraph_format.left_indent = Cm(0.45)
    font("CV Tools", size=7.2, before=3, after=10, line=1.9)
    quote = font("CV Quote", size=9, color="3A3A3A", after=3)
    quote.font.italic = True
    quote.paragraph_format.left_indent = Cm(0.5)
    border(quote, "left", RED, 18, 8)
    font("Sidebar Sections Spacing", size=1, after=14.8, line=1.0)
    font("Main Sections Spacing", size=1, after=14.8, line=1.0)
    char_style(doc, "CV Accent", font=FONT, color=RED)
    char_style(doc, "CV Muted", font=FONT, color=MUTED, size=8)
    tool = char_style(doc, "CV Tool", font=FONT, color=MUTED, size=6.4)
    tool.font.all_caps = True
    tool.element.get_or_add_rPr().append(xml(f'<w:bdr w:val="single" w:sz="6" w:space="0" w:color="{LINE}"/>'))
    char_style(doc, "CV Tool Mark", font=FONT, size=6.4)
    char_style(doc, "CV Skill Fill", font=FONT, fill=RED)
    char_style(doc, "CV Skill Track", font=FONT, fill=LINE)
    for s in ("Heading 1", "Heading 2", "Heading 3", "Heading 4", "CV Role", "CV Company", "Date"):
        doc.styles[s].paragraph_format.keep_with_next = True


def header(doc: Doc) -> None:
    """Header rows of fixed height, so the body starts right under the headline numbers."""
    head = table(doc, [(M + NAME_W) / 10, (ROLE_X - M - NAME_W) / 10, (210 - ROLE_X) / 10], rows=3)
    for row, mm in zip(head.rows, (CONTACT, NAME, STATS), strict=True):
        row.height, row.height_rule = Cm(mm / 10), WD_ROW_HEIGHT_RULE.EXACTLY
    top = head.rows[0].cells[0].merge(head.rows[0].cells[2])
    cell(top, 21)
    top.paragraphs[0]._p.style = "CVPhoto"
    name_c, _, role_c = head.rows[1].cells
    cell(name_c, (M + NAME_W) / 10, pad=(0, 0, 0.87, M / 10), valign=WD_CELL_VERTICAL_ALIGNMENT.BOTTOM)
    cell(role_c, (210 - ROLE_X) / 10, pad=(1.4, M / 10, 0, 0))
    cell(head.rows[1].cells[1], (ROLE_X - M - NAME_W) / 10)
    name_c.paragraphs[0]._p.style = "Name"
    name_c.paragraphs[0].add_run("{Name}")
    role_c.paragraphs[0]._p.style = "JobTitle"
    role_c.paragraphs[0].add_run("{Job Title}")
    stats = head.rows[2].cells[0].merge(head.rows[2].cells[2])
    cell(stats, 21, pad=(0.35, M / 10, 0, M / 10))
    stats.paragraphs[0]._p.style = "CVHighlights"
    stats.paragraphs[0].add_run("{Highlights}")


def layout(doc: Doc) -> None:
    """Body: sidebar column and main column."""
    doc.add_paragraph(style="CV Photo")  # keeps Word from merging the two tables
    side, main = table(doc, [(M + SIDE_W) / 10, 21 - (M + SIDE_W) / 10]).rows[0].cells
    cell(side, (M + SIDE_W) / 10, pad=(0, 0, 0, M / 10))
    cell(main, 21 - (M + SIDE_W) / 10, pad=(0, M / 10, 0, GAP / 10))
    side.paragraphs[0]._p.style = "CVPhoto"  # word.py adds a section spacer of its own
    doc.add_paragraph(style="CV Photo")  # Word needs a paragraph between the last table and the section end


def grid() -> list[tuple[float, float, float, float, str]]:
    """The 12-column hairlines, as the CSS draws them 13 mm in from each edge."""
    step = (WIDTH + 0.26) / 12
    return [(M + k * step, 0, 0.26, 297, HAIR) for k in range(13)]


def decorate(doc: Doc) -> None:
    """Page decoration as shapes behind the text, repeated by the page headers: the column hairlines, and on
    page 1 the header's rules, red bars and column dividers. Later pages start the body 14 mm down."""
    sec = doc.sections[0]
    first = sec.first_page_header.paragraphs[0]
    first._p.style = "CVPhoto"
    rules = [
        (M, TOP + CONTACT - 0.26, WIDTH, 0.26, INK),
        (M, STATS_Y, WIDTH, 2, RED),
        (M, STATS_Y + 5, WIDTH, 0.26, INK),
        (ROLE_X, 30, 2.4, 2.4, RED),
        *[(M + k * WIDTH / 4 - 2.6, STATS_Y + 7.5, 0.26, BODY_Y - STATS_Y - 7.5, LINE) for k in (1, 2, 3)],
    ]
    for i, box in enumerate([*grid(), *rules]):
        rect(first, box, 1001 + i)
    later = sec.header.paragraphs[0]
    later._p.style = "CVPhoto"
    later.paragraph_format.space_before = Cm(1.32)
    for i, box in enumerate(grid()):
        rect(later, box, 1101 + i)


def make(out: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(0)
    sec.top_margin, sec.bottom_margin = Cm(TOP / 10), Cm(1.2)
    sec.header_distance, sec.footer_distance = Cm(0), Cm(0)
    sec.different_first_page_header_footer = True
    styles(doc)
    header(doc)
    layout(doc)
    decorate(doc)
    finish(doc, out, FONT)
