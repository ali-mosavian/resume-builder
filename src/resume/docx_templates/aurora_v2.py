"""templates/aurora-v2.docx, the Word twin of themes/aurora-v2.css.

Dark gradient header with headline numbers, sidebar with skill bars, role timeline, tool chips and quote bars.
"""

from typing import Any
from pathlib import Path

from docx import Document
from docx.shared import Cm
from docx.document import Document as Doc
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.styles.style import ParagraphStyle
from docx.enum.table import WD_ROW_HEIGHT_RULE

from resume.ooxml import xml
from resume.docx_templates import Mark
from resume.docx_templates import cell
from resume.docx_templates import rect
from resume.docx_templates import style
from resume.docx_templates import table
from resume.docx_templates import border
from resume.docx_templates import finish
from resume.docx_templates import numbered
from resume.docx_templates import char_style
from resume.docx_templates import bullet_lists

SERIF, SANS = "Fraunces", "Inter"
INK, MUTED, LINE, VIOLET, CHIP, CHIP_INK = "151A2D", "667085", "E4E7EF", "7C5CFF", "EEF2FF", "3730A3"
GAPS = {
    SERIF: 1.233,
    SANS: 1.21,
}  # hhea line height per em: Word's "multiple" spacing scales it, CSS line-height doesn't

# Page geometry in mm, as themes/aurora-v2.css lays out an A4 sheet.
TOP, HEAD1, HEAD2, SIDE_W = 10, 33, 19, 52  # header block: name rows, then headline numbers; sidebar width
LATER_TOP = 12  # where the body starts on later pages


def styles(doc: Doc) -> None:
    def font(
        name: str, *, face: str = SANS, base: str | None = "Normal", color: str = INK, line: float = 1.42, **kw: Any
    ) -> ParagraphStyle:
        return style(doc, name, base=base, font=face, color=color, line=line, line_gap=GAPS[face], **kw)

    _, _dot, ring = bullet_lists(
        doc,
        [
            Mark("•", VIOLET, 652, 255, 10, "Segoe UI Symbol"),
            Mark("•", VIOLET, 652, 255, 10, "Segoe UI Symbol"),
            Mark("○", VIOLET, 397, 397, 11, "Segoe UI Symbol"),
        ],
    )
    font("Normal", base=None, size=9.2, after=7)
    font("Name", face=SERIF, size=30, bold=True, color="FFFFFF", after=0, line=1.02)
    font("Job Title", size=7.6, color="8FD0E6", caps=True, spacing=1.2, before=8, after=0)
    font("CV Photo", size=2, after=0, line=1.0)
    font("CV Highlights", size=2, after=0)
    font("CV Highlight Value", face=SERIF, size=15, bold=True, color="E0E7FF", after=0, line=1.1)
    font("CV Highlight Label", size=7.2, color="A5B4FC", after=0)

    font("Heading 1", face=SERIF, size=14, bold=True, after=8.5, line=1.2)
    font("Heading 2", size=9.5, bold=True, before=0, after=0)
    numbered(font("CV Role", face=SERIF, size=10.5, bold=True, after=0, line=1.3), ring)
    font("CV Company", size=8.2, color=MUTED, after=2).paragraph_format.left_indent = Cm(0.7)
    font("Date", size=7, color=MUTED, caps=True, spacing=0.7, before=2, after=7)
    role_date = font("CV Role Date", base="Date", size=7, color=MUTED, caps=True, spacing=0.7, before=2, after=7)
    role_date.paragraph_format.left_indent = Cm(0.7)
    font("Heading 3", size=6.6, bold=True, color=MUTED, caps=True, spacing=0.9, after=5.7, line=1.3)
    font("Heading 4", size=6.2, bold=True, color=MUTED, caps=True, spacing=0.8, before=5.7, after=0)
    font("Sidebar Text", size=8.2, after=1.5)
    kv = font("CV Key Value", size=8.2, after=4.5)
    kv.paragraph_format.tab_stops.add_tab_stop(Cm(SIDE_W / 10), alignment=WD_TAB_ALIGNMENT.RIGHT)
    font("CV Skill", size=8.2, after=1.5)
    font("CV Skill Bar", size=1, after=5.5, line=1.0)
    list_p = font("List Paragraph", size=9.2, after=2.3)
    list_p.paragraph_format.left_indent = Cm(1.15)
    font("CV Tools", size=7, after=11, line=1.8).paragraph_format.left_indent = Cm(0.7)
    quote = font("CV Quote", size=9.2, color="374151", after=4.3)
    quote.paragraph_format.left_indent = Cm(0.3)
    border(quote, "left", "A5B4FC", 18, 5)
    font("Sidebar Sections Spacing", size=4, after=12)
    font("Main Sections Spacing", size=4, after=15)
    char_style(doc, "CV Accent", font=SANS, color=VIOLET)
    char_style(doc, "CV Muted", font=SANS, color=MUTED, size=7.5)
    char_style(doc, "CV Tool", font=SANS, fill=CHIP, color=CHIP_INK, size=7)
    char_style(doc, "CV Tool Mark", font=SANS, fill=CHIP, size=7)
    char_style(doc, "CV Skill Fill", font=SANS, fill=VIOLET)
    char_style(doc, "CV Skill Track", font=SANS, fill=LINE)
    for s in ("CV Role", "CV Company", "Date", "CV Role Date", "Heading 1", "Heading 3", "Heading 4"):
        doc.styles[s].paragraph_format.keep_with_next = True


def header(doc: Doc) -> None:
    """Two rows of fixed height over the header shape: name and role, then the headline numbers."""
    head = table(doc, [21.0], rows=2)
    for row, mm in zip(head.rows, (HEAD1, HEAD2), strict=True):
        row.height, row.height_rule = Cm(mm / 10), WD_ROW_HEIGHT_RULE.EXACTLY
    name_c, stats = head.rows[0].cells[0], head.rows[1].cells[0]
    cell(name_c, 21.0, pad=(0.8, 2.1, 0, 2.1))
    cell(stats, 21.0, pad=(0.4, 2.1, 0, 2.1))
    name_c.paragraphs[0]._p.style = "Name"
    name_c.paragraphs[0].add_run("{Name}")
    name_c.add_paragraph(style="Job Title").add_run("{Job Title}")
    stats.paragraphs[0]._p.style = "CVHighlights"
    stats.paragraphs[0].add_run("{Highlights}")


def layout(doc: Doc) -> None:
    doc.add_paragraph(style="CV Photo")  # keeps Word from merging the two tables
    side_w = 1.2 + SIDE_W / 10 + 0.6
    side, main = table(doc, [side_w, 21 - side_w]).rows[0].cells
    cell(side, side_w, pad=(0, 0.6, 0, 1.2))
    cell(main, 21 - side_w, pad=(0.5, 1.2, 0, 0))
    side.paragraphs[0]._p.style = "CVPhoto"
    doc.add_paragraph(style="CV Photo")


def shape(
    p: Paragraph, box: tuple[float, float, float, float], shape_id: int, fill: str, *, prst: str = "rect", adj: str = ""
) -> None:
    """rect() with another fill (gradFill xml) and preset geometry."""
    rect(p, (*box, "000000"), shape_id)
    r = p._p[-1]
    solid = r.xpath(".//a:solidFill")[0]
    solid.getparent().replace(solid, xml(fill, "a"))
    geom = r.xpath(".//a:prstGeom")[0]
    geom.set("prst", prst)
    if adj:
        geom.getparent().replace(geom, xml(f'<a:prstGeom prst="{prst}"><a:avLst>{adj}</a:avLst></a:prstGeom>', "a"))


def decorate(doc: Doc) -> None:
    sec = doc.sections[0]
    first = sec.first_page_header.paragraphs[0]
    first._p.style = "CVPhoto"
    grad = (
        '<a:gradFill rotWithShape="1"><a:gsLst><a:gs pos="0"><a:srgbClr val="0B1024"/></a:gs>'
        '<a:gs pos="100000"><a:srgbClr val="1A2150"/></a:gs></a:gsLst><a:lin ang="2700000" scaled="0"/></a:gradFill>'
    )
    # 4 mm corners of a 52 mm tall shape: adj is a fraction of the shorter side.
    shape(first, (12, TOP, 186, HEAD1 + HEAD2), 1001, grad, prst="roundRect", adj='<a:gd name="adj" fmla="val 7692"/>')
    shape(first, (21, TOP + HEAD1, 168, 0.3), 1002, '<a:solidFill><a:srgbClr val="2B3150"/></a:solidFill>')
    later = sec.header.paragraphs[0]
    later._p.style = "CVPhoto"
    later.paragraph_format.space_before = Cm((LATER_TOP - 5) / 10)


def make(out: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(0)
    sec.top_margin, sec.bottom_margin = Cm(TOP / 10), Cm(1.2)
    sec.header_distance, sec.footer_distance = Cm(0), Cm(0)
    sec.different_first_page_header_footer = True
    styles(doc)
    header(doc)
    layout(doc)
    decorate(doc)
    finish(doc, out, SERIF, SANS)
