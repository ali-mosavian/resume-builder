"""templates/bracket.docx, the Word twin of themes/bracket.css.

Beyond the style ids every template has, it opts into word.py's optional parts: Photo, Contact and
Highlights placeholders, skill bars, a role timeline, tool chips and quote bars.
"""

import uuid
from functools import partial
from pathlib import Path

from docx import Document
from docx.document import Document as Doc
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_UNDERLINE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.oxml import serialize_part_xml
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.styles.style import ParagraphStyle
from docx.text.paragraph import Paragraph

from resume import ROOT
from resume.docx_templates import Mark, border, bullet_lists, cell, char_style, numbered, style, table
from resume.ooxml import xml

FONT = "Montserrat"
INK, MUTED, ACCENT, SIDE, TRACK = "262626", "7A7A7A", "F6B92B", "F4F3EF", "DDDBD3"
LINE_GAP = 1.219  # Montserrat's line height per em: Word's "multiple" spacing scales it, CSS line-height doesn't

# Page geometry in mm, as themes/bracket.css lays out an A4 sheet.
TOP, HEAD1, HEAD2, SIDE_W = 7, 45, 24, 64  # header rows: name block, then headline numbers; sidebar width
SIDE_Y = TOP + HEAD1 + HEAD2  # where the sidebar starts on page 1


def accent_underline(s: ParagraphStyle) -> None:
    s.font.underline = WD_UNDERLINE.THICK
    s.element.get_or_add_rPr().xpath("./w:u")[0].set(qn("w:color"), ACCENT)


def embed_font(doc: Doc, family: str, files: dict[str, Path]) -> None:
    """Embed TrueType files (keys: Regular, Bold, Italic, BoldItalic) so Word shows the font without it installed.

    Word stores each as an obfuscated .odttf: the first 32 bytes XORed with its GUID key, bytes reversed.
    """
    font_table = doc.part.part_related_by(RT.FONT_TABLE)
    root = parse_xml(font_table.blob)
    for old in root.findall(qn("w:font")):  # a font table isn't an oxml class: no w: prefix in xpath
        if old.get(qn("w:name")) == family:
            root.remove(old)
    font = xml(f'<w:font w:name="{family}"><w:charset w:val="00"/><w:family w:val="swiss"/><w:pitch w:val="variable"/></w:font>')
    for i, (kind, path) in enumerate(files.items(), 1):
        guid = str(uuid.uuid4()).upper()
        key = bytes.fromhex(guid.replace("-", ""))[::-1]
        data = bytearray(path.read_bytes())
        for j in range(32):
            data[j] ^= key[j % 16]
        part = Part(
            PackURI(f"/word/fonts/font{i}.odttf"),
            "application/vnd.openxmlformats-officedocument.obfuscatedFont",
            bytes(data),
            doc.part.package,
        )
        font.append(xml(f'<w:embed{kind} r:id="{font_table.relate_to(part, RT.FONT)}" w:fontKey="{{{guid}}}"/>', "w", "r"))
    root.append(font)
    font_table._blob = serialize_part_xml(root)
    doc.settings.element.xpath("./w:zoom")[0].addnext(xml("<w:embedTrueTypeFonts/>"))  # schema: right after zoom


WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"


type Box = tuple[float, float, float, float, str]  # x, y, width, height in mm; colour


def rect(p: Paragraph, box: Box, shape_id: int) -> None:
    """A filled rectangle behind the text, placed on the page: what CSS draws with positioned boxes."""
    x, y, w, h, color = box

    def emu(mm: float) -> int:
        return int(mm * 36000)

    p._p.append(
        xml(
            f'<w:r xmlns:wps="{WPS}"><w:drawing>'
            f'<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="{shape_id}" behindDoc="1" '
            f'locked="1" layoutInCell="0" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
            f'<wp:positionH relativeFrom="page"><wp:posOffset>{emu(x)}</wp:posOffset></wp:positionH>'
            f'<wp:positionV relativeFrom="page"><wp:posOffset>{emu(y)}</wp:posOffset></wp:positionV>'
            f'<wp:extent cx="{emu(w)}" cy="{emu(h)}"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
            f'<wp:docPr id="{shape_id}" name="Shape {shape_id}"/><wp:cNvGraphicFramePr/>'
            f'<a:graphic><a:graphicData uri="{WPS}"><wps:wsp><wps:cNvSpPr/><wps:spPr>'
            f'<a:xfrm><a:off x="0" y="0"/><a:ext cx="{emu(w)}" cy="{emu(h)}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
            f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill><a:ln><a:noFill/></a:ln></wps:spPr><wps:bodyPr/></wps:wsp>'
            f"</a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>",
            "w",
            "wp",
            "a",
        )
    )


def styles(doc: Doc) -> None:
    font = partial(style, doc, base="Normal", font=FONT, color=INK, line=1.6, line_gap=LINE_GAP)

    # One list per style: Word links a list to a single style and drops the marks of any other.
    _, diamond, diamond_side, ring = bullet_lists(  # word.py's bullets take the first
        doc,
        [
            Mark("•", ACCENT, 624, 284, 10, "Segoe UI Symbol"),
            Mark("◆", ACCENT, 400, 400, 11, "Segoe UI Symbol"),
            Mark("◆", ACCENT, 400, 400, 11, "Segoe UI Symbol"),
            Mark("○", INK, 340, 340, 12, "Segoe UI Symbol"),
        ],
    )

    font("Normal", base=None, size=9.5, after=4)
    font("Name", size=27.5, bold=True, caps=True, after=0, line=1.0)
    font("Job Title", size=6.8, color=MUTED, caps=True, spacing=1.3, after=3)
    font("CV Photo", size=2, after=0, line=1.0)
    contact = font("CV Contact", size=8, bold=True, before=0, after=0, line=1.0)
    contact.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    border(contact, "top between bottom", INK, 6, 5)
    font("CV Highlights", size=2, after=0)
    accent_underline(font("CV Highlight Value", size=18, bold=True, after=4, line=1.0))
    font("CV Highlight Label", size=7.4, color=MUTED, after=0)

    numbered(font("Heading 1", size=12.8, bold=True, before=0, after=8, line=1.2), diamond)
    font("Heading 2", size=9.7, bold=True, before=4, after=0)
    numbered(font("CV Role", base="Heading 2", size=9.7, bold=True, before=8, after=0), ring)
    font("CV Company", size=8.2, color=MUTED, after=1).paragraph_format.left_indent = Cm(0.6)
    accent_underline(font("Date", size=7.2, bold=True, spacing=0.3, after=6))
    role_date = font("CV Role Date", base="Date", size=7.2, bold=True, spacing=0.3, after=6)
    accent_underline(role_date)
    role_date.paragraph_format.left_indent = Cm(0.6)  # under the role title, past the ring
    numbered(font("Heading 3", size=10.6, bold=True, before=0, after=6, line=1.2), diamond_side)
    font("Heading 4", size=6.6, color=MUTED, bold=True, caps=True, spacing=1.1, before=6, after=0)
    font("Sidebar Text", size=8.7, after=1)
    kv = font("CV Key Value", size=8.7, after=3)
    kv.paragraph_format.tab_stops.add_tab_stop(Cm(4.6), alignment=WD_TAB_ALIGNMENT.RIGHT)  # at the sidebar text edge
    font("CV Skill", size=8.7, after=2)
    font("CV Skill Bar", size=1, after=6, line=1.0)
    font("List Paragraph", size=9.5, after=2).paragraph_format.left_indent = Cm(0.6)
    font("CV Tools", size=7.2, after=10, line=1.8).paragraph_format.left_indent = Cm(0.6)
    quote = font("CV Quote", size=9.5, color="3A3A3A", after=3)
    quote.font.italic = True
    quote.paragraph_format.left_indent = Cm(0.4)
    border(quote, "left", ACCENT, 18, 8)
    font("Sidebar Sections Spacing", size=4, after=16)
    font("Main Sections Spacing", size=4, after=26)
    char_style(doc, "CV Accent", font=FONT, color=ACCENT)
    char_style(doc, "CV Muted", font=FONT, color=MUTED, size=7.5)
    char_style(doc, "CV Tool", font=FONT, fill=SIDE, size=7.2)
    char_style(doc, "CV Tool Mark", font=FONT, fill=ACCENT, size=7.2)  # a chip's yellow left edge
    char_style(doc, "CV Skill Fill", font=FONT, fill=ACCENT)
    char_style(doc, "CV Skill Track", font=FONT, fill=TRACK)
    for s in ("CV Contact", "CV Role", "CV Company", "Date", "CV Role Date", "Heading 1", "Heading 3"):
        doc.styles[s].paragraph_format.keep_with_next = True


def header(doc: Doc) -> None:
    """Header rows of fixed height, so page 1's sidebar shape starts right under the headline numbers."""
    head = table(doc, [5.7, 9.3, 6.0], rows=2)
    for row, mm in zip(head.rows, (HEAD1, HEAD2), strict=True):
        row.height, row.height_rule = Cm(mm / 10), WD_ROW_HEIGHT_RULE.EXACTLY
    photo_c, name_c, contact_c = head.rows[0].cells
    cell(photo_c, 5.7, pad=(0.9, 0, 0, 1.6))
    cell(name_c, 9.3, pad=(0.9, 0.4, 0, 0.6), valign=WD_CELL_VERTICAL_ALIGNMENT.CENTER)
    cell(contact_c, 6.0, pad=(0.9, 1.4, 0, 0.4), valign=WD_CELL_VERTICAL_ALIGNMENT.CENTER)
    photo_c.paragraphs[0]._p.style = "CVPhoto"
    photo_c.paragraphs[0].add_run("{Photo 3.1 gray}")
    name_c.paragraphs[0]._p.style = "JobTitle"
    name_c.paragraphs[0].add_run("{Job Title}")
    name_c.add_paragraph(style="Name").add_run("{Name}")
    contact_c.paragraphs[0]._p.style = "CVContact"
    contact_c.paragraphs[0].add_run("{Contact}")
    stats = head.rows[1].cells[0].merge(head.rows[1].cells[2])
    cell(stats, 21, pad=(0.4, 1.4, 0, 1.6))
    stats.paragraphs[0]._p.style = "CVHighlights"
    stats.paragraphs[0].add_run("{Highlights}")


def layout(doc: Doc) -> None:
    """Body: sidebar column (its gray is a shape, see decorate) and main column."""
    doc.add_paragraph(style="CV Photo")  # keeps Word from merging the two tables
    side, main = table(doc, [SIDE_W / 10, 21 - SIDE_W / 10]).rows[0].cells
    cell(side, SIDE_W / 10, pad=(0.3, 0.8, 0, 1.0))
    cell(main, 21 - SIDE_W / 10, pad=(1.0, 1.4, 0, 1.0))
    side.paragraphs[0]._p.style = "CVPhoto"  # word.py adds a section spacer of its own
    doc.add_paragraph(style="CV Photo")  # Word needs a paragraph between the last table and the section end


def decorate(doc: Doc) -> None:
    """Page decoration as shapes behind the text, repeated by the page headers. Page 1: the bracket, its
    black square, the photo's offset block, the sidebar from under the header down. Later pages: the
    sidebar from the top, and the body starts 14 mm down, as the web's later sheets do."""
    sec = doc.sections[0]
    first = sec.first_page_header.paragraphs[0]
    first._p.style = "CVPhoto"
    for i, box in enumerate(
        [
            (7, 7, 60, 5, ACCENT),
            (7, 7, 5, 30, ACCENT),
            (14, 12, 2.2, 2.2, INK),
            (19, 19, 31, 31, ACCENT),
            (0, SIDE_Y, SIDE_W, 297 - SIDE_Y, SIDE),
            (188, 292, 22, 5, ACCENT),
        ]
    ):
        rect(first, box, 1001 + i)
    later = sec.header.paragraphs[0]
    later._p.style = "CVPhoto"
    later.paragraph_format.space_before = Cm(1.3)
    for i, box in enumerate([(0, 0, SIDE_W, 297, SIDE), (188, 292, 22, 5, ACCENT)]):
        rect(later, box, 1101 + i)


def make(out: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(0)
    sec.top_margin, sec.bottom_margin = Cm(TOP / 10), Cm(1.2)
    sec.header_distance, sec.footer_distance = Cm(0), Cm(0)
    sec.different_first_page_header_footer = True
    styles(doc)
    header(doc)
    layout(doc)
    decorate(doc)
    settings = doc.settings.element
    settings.xpath("./w:zoom")[0].set(qn("w:percent"), "100")  # python-docx's blank document omits this required attribute
    # Word 2013+ layout. In the blank document's Word 2010 mode, a table's first-cell text aligns with the margin,
    # so the cell padding pushes the table, the sidebar text and its marks off the page's left edge.
    settings.xpath("./w:compat/w:compatSetting[@w:name='compatibilityMode']")[0].set(qn("w:val"), "15")
    fonts = ROOT / "assets" / "fonts"
    embed_font(doc, FONT, {kind: fonts / f"{FONT}-{kind}.ttf" for kind in ("Regular", "Bold", "Italic", "BoldItalic")})
    doc.save(str(out))
