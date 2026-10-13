"""templates/cinematic.docx, the Word twin of themes/cinematic.css (its A4 sheets, at rest: no animation).

Dark page drawn as shapes in the page headers; role timeline, skill bars, tool chips and quote bars as in bracket.py.
"""

from pathlib import Path
from functools import partial

from docx import Document
from docx.shared import Cm
from docx.document import Document as Doc
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.styles.style import ParagraphStyle
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ROW_HEIGHT_RULE

from resume.ooxml import xml
from resume.docx_templates import Mark
from resume.docx_templates import cell
from resume.docx_templates import rect
from resume.docx_templates import style
from resume.docx_templates import table
from resume.docx_templates import border
from resume.docx_templates import finish
from resume.docx_templates import numbered
from resume.docx_templates import char_style
from resume.docx_templates import bullet_lists

SANS, SERIF, MONO = "Instrument Sans", "Instrument Serif", "JetBrains Mono"
BG, INK, INK2, MUTED, AMBER = "0B0B10", "F3EFE7", "D6D0C4", "A39D92", "F0B35A"
PANEL, SIDE, LINE, CHIP, TRACK = "0F0F16", "12121A", "25252C", "1C1C25", "25252E"
# Line height per em of each font (Word's "multiple" spacing scales it, CSS line-height doesn't).
GAP = {SANS: 1.22, SERIF: 1.3, MONO: 1.32}

# Page geometry in mm, as themes/cinematic.css lays out an A4 sheet.
TOP, BOTTOM, HEAD1, HEAD2, SIDE_W = 8, 10, 40, 34, 58  # header rows: name block, then headline numbers
SIDE_Y = TOP + HEAD1 + HEAD2  # where the sidebar starts on page 1
GUTTER = 14  # page side padding of the header
STATS = (GUTTER, 50, 210 - 2 * GUTTER, 26)  # the headline panel: x, y, w, h


def pt(mm: float) -> float:
    return mm * 72 / 25.4


def styles(doc: Doc) -> None:
    def font(
        name: str,
        *,
        face: str = SANS,
        size: float,
        line: float = 1.5,
        line_pt: float | None = None,
        color: str = INK2,
        bold: bool = False,
        caps: bool = False,
        spacing: float = 0,
        before: float = 0,
        after: float = 2,
    ) -> ParagraphStyle:
        """line_pt: the CSS line box in points; otherwise line is a multiple of the font size."""
        return style(
            doc,
            name,
            "Normal",
            font=face,
            size=size,
            color=color,
            bold=bold,
            caps=caps,
            spacing=spacing,
            before=before,
            after=after,
            line=line if line_pt is None else line_pt / size,
            line_gap=GAP[face],
        )

    mark = partial(Mark, font="Segoe UI Symbol")
    _, ring = bullet_lists(doc, [mark("•", AMBER, 624, 200, 10), mark("○", AMBER, 340, 340, 12)])
    numbering = doc.part.numbering_part.element
    numbering.insert(
        2,
        xml(
            '<w:abstractNum w:abstractNumId="95"><w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimalZero"/>'
            '<w:lvlText w:val="%1"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="510" w:hanging="510"/></w:pPr>'
            f'<w:rPr><w:rFonts w:ascii="{MONO}" w:hAnsi="{MONO}"/><w:color w:val="{AMBER}"/><w:sz w:val="14"/></w:rPr></w:lvl></w:abstractNum>'
        ),
    )
    numbering.append(xml('<w:num w:numId="95"><w:abstractNumId w:val="95"/></w:num>'))

    style(doc, "Normal", None, font=SANS, size=9, color=INK2, after=pt(3), line=1.5, line_gap=GAP[SANS])
    font("Name", face=SERIF, size=46, color=INK, after=0, line=1.0, spacing=-1.6)
    font("Job Title", face=MONO, size=6.8, color=AMBER, caps=True, spacing=1.4, after=pt(1.5), line_pt=6.8 * 1.4)
    font("CV Photo", size=1, after=0, line=1.0)
    contact = font("CV Contact", size=8, color=INK, bold=True, before=0, after=0, line_pt=11.2)
    contact.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    border(contact, "top between bottom", LINE, 2, 5)
    font("CV Highlights", size=1, after=0, line=1.0)
    font("CV Highlight Value", face=SERIF, size=24, color=INK, after=pt(1), line=1.0)
    font("CV Highlight Label", size=7.4, color=MUTED, after=0)

    numbered(font("Heading 1", face=SERIF, size=16, color=INK, before=pt(10), after=pt(4), line=1.2), "95")
    font("Heading 2", size=9.7, bold=True, color=INK, before=0, after=0)
    numbered(font("CV Role", face=SERIF, size=12, color=INK, before=0, after=pt(0.5), line=1.25), ring)
    font("CV Company", size=8.2, color=MUTED, after=0, line_pt=14.55).paragraph_format.left_indent = Cm(0.6)
    font("Date", face=MONO, size=7, color=AMBER, spacing=0.4, before=pt(0.6), after=pt(2), line=1.4)
    role_date = font("CV Role Date", face=MONO, size=7, color=AMBER, spacing=0.4, before=pt(0.6), after=pt(2), line=1.4)
    role_date.paragraph_format.left_indent = Cm(0.6)
    font("Heading 3", face=MONO, size=6.6, color=AMBER, caps=True, spacing=1.3, before=0, after=pt(3), line=1.0)
    font("Heading 4", face=MONO, size=6, color=MUTED, caps=True, spacing=0.8, before=pt(2.4), after=pt(0.3), line=1.32)
    font("Sidebar Text", size=8.7, after=0)
    kv = font("CV Key Value", size=8.7, after=pt(1.6))
    kv.paragraph_format.tab_stops.add_tab_stop(Cm(4.0), alignment=WD_TAB_ALIGNMENT.RIGHT)
    font("CV Skill", size=8.7, color=INK, after=pt(1))
    font("CV Skill Bar", size=1, after=pt(2.4), line=1.0)
    font("List Paragraph", size=9, after=pt(0.6)).paragraph_format.left_indent = Cm(0.6)
    font("CV Tools", size=7.2, before=pt(2.4), after=pt(4), line=2.1).paragraph_format.left_indent = Cm(0.6)
    quote = font("CV Quote", face=SERIF, size=11, after=pt(1.5), line=1.4)
    quote.font.italic = True
    quote.paragraph_format.left_indent = Cm(0.35)
    border(quote, "left", AMBER, 18, 8)
    font("Sidebar Sections Spacing", size=1, after=pt(6) - 1, line=1.0)
    font("Main Sections Spacing", size=1, after=0, line=1.0)
    char_style(doc, "CV Accent", font=SANS, color=AMBER)
    char_style(doc, "CV Muted", font=SANS, color=MUTED, size=7.5)
    char_style(doc, "CV Tool", font=SANS, fill=CHIP, size=7.2)
    char_style(doc, "CV Tool Mark", font=SANS, fill=AMBER, size=7.2)  # a chip's amber left edge
    char_style(doc, "CV Skill Fill", font=SANS, fill=AMBER)
    char_style(doc, "CV Skill Track", font=SANS, fill=TRACK)
    for s in (
        "CV Contact",
        "CV Role",
        "CV Company",
        "Date",
        "CV Role Date",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "Heading 4",
    ):
        doc.styles[s].paragraph_format.keep_with_next = True


def header(doc: Doc) -> None:
    """Header rows of fixed height, so page 1's sidebar shape starts right under the headline numbers."""
    right = 50  # contact column width
    head = table(doc, [21 - 0.8 - 1.4 - right / 10, right / 10 + 0.8 + 1.4], rows=2)
    for row, mm in zip(head.rows, (HEAD1, HEAD2), strict=True):
        row.height, row.height_rule = Cm(mm / 10), WD_ROW_HEIGHT_RULE.EXACTLY
    name_c, contact_c = head.rows[0].cells
    cell(name_c, head.columns[0].width.cm, pad=(1.51, 0, 0, 1.4))
    cell(
        contact_c, head.columns[1].width.cm, pad=(0, 1.4, 0, 0.8)
    )  # Word sits the links on the row's bottom edge, whatever the padding
    name_c.paragraphs[0]._p.style = "JobTitle"
    name_c.paragraphs[0].add_run("{Job Title}")
    name_c.add_paragraph(style="Name").add_run("{Name}")
    contact_c.paragraphs[0]._p.style = "CVContact"
    contact_c.paragraphs[0].add_run("{Contact}")
    stats = head.rows[1].cells[0].merge(head.rows[1].cells[1])
    cell(stats, 21, pad=(0.725, 1.215, 0, 1.635))
    stats.paragraphs[0]._p.style = "CVHighlights"
    stats.paragraphs[0].add_run("{Highlights}")


def layout(doc: Doc) -> None:
    """Body: sidebar column (its panel is a shape, see decorate) and main column."""
    doc.add_paragraph(style="CV Photo")  # keeps Word from merging the two tables
    side, main = table(doc, [SIDE_W / 10, 21 - SIDE_W / 10]).rows[0].cells
    cell(side, SIDE_W / 10, pad=(0, 0.8, 0, 1.0))
    cell(main, 21 - SIDE_W / 10, pad=(0, 1.4, 0, 1.0))
    side.paragraphs[0]._p.style = "CVPhoto"  # word.py adds a section spacer of its own
    doc.add_paragraph(style="CV Photo")  # Word needs a paragraph between the last table and the section end


def decorate(doc: Doc) -> None:
    """The dark page as shapes behind the text, repeated by the page headers: background, letterbox bars,
    sidebar panel and the headline panel (page 1). Later pages: sidebar from the top; the body starts 14 mm down."""
    sec = doc.sections[0]
    sx, sy, sw, sh = STATS
    bars = [(0, 0, 210, TOP, "000000"), (0, TOP, 210, 0.25, LINE), (0, 291, 210, 6, "000000")]
    first = sec.first_page_header.paragraphs[0]
    first._p.style = "CVPhoto"
    page1 = [
        (0, 0, 210, 297, BG),
        (0, SIDE_Y, SIDE_W, 297 - SIDE_Y, SIDE),
        (sx, sy, sw, sh, LINE),
        (sx + 0.25, sy + 0.25, sw - 0.5, sh - 0.5, PANEL),
        *bars,
    ]
    for i, box in enumerate(page1):
        rect(first, box, 1001 + i)
    later = sec.header.paragraphs[0]
    later._p.style = "CVPhoto"
    later.paragraph_format.space_before = Cm(1.38)
    for i, box in enumerate([(0, 0, 210, 297, BG), (0, 0, SIDE_W, 297, SIDE), *bars]):
        rect(later, box, 1101 + i)


def make(out: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(0)
    sec.top_margin, sec.bottom_margin = Cm(TOP / 10), Cm(BOTTOM / 10)
    sec.header_distance, sec.footer_distance = Cm(0), Cm(0)
    sec.different_first_page_header_footer = True
    styles(doc)
    header(doc)
    layout(doc)
    decorate(doc)
    finish(doc, out, SANS, SERIF, MONO)
