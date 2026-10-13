"""Build a CV .docx from the YAML, styled by a .docx template.

The template's last top-level table is the layout: first column the sidebar, last column the main body.
Restyle the template in Word; paragraph styles are referenced by style id below.

Optional parts, used when the template has them (see docx_templates/bracket.py):
placeholders CVPhoto ("{Photo <width cm> [gray]}"), CVContact and CVHighlights; styles CVRole + CVCompany
(role timeline), CVTool + CVTools (tool chips, no labels), CVSkill + CVSkillBar (level bars), CVKeyValue, CVQuote.
The CV prefix keeps them from matching a template's own styles (Word has a built-in Quote).
"""

from pathlib import Path
from collections.abc import Iterable
from collections.abc import Iterator

from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn
from docx.shared import Emu
from docx.table import Table
from docx.table import _Cell
from docx.shared import Twips
from docx.text.run import Run
from docx.shared import Length
from docx.oxml import OxmlElement
from docx.oxml.styles import CT_Style
from docx.document import Document as Doc
from docx.text.paragraph import Paragraph
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.opc.constants import RELATIONSHIP_TYPE

from resume.data import CV
from resume.data import BOLD
from resume.data import span
from resume.data import when
from resume.ooxml import xml
from resume.data import Labels
from resume.ooxml import shading


def add_text(p: Paragraph, text: str) -> None:
    """Append text with **bold** markup as runs: in the template's CVStrong style if it has one, else bold."""
    strong = has(p, "CVStrong")
    for i, part in enumerate(BOLD.split(text)):
        if not part:
            continue
        run = p.add_run(part)
        if i % 2 and strong:
            styled(run, "CVStrong")
        else:
            # None, not False, outside markup: False would override bold paragraph styles.
            run.bold = True if i % 2 else None


def para(cell: _Cell, style_id: str, text: str = "") -> Paragraph:
    # By style id, not name: the default template has duplicate style names.
    p = cell.add_paragraph()
    p._p.style = style_id
    add_text(p, text)
    return p


def style(block: _Cell | Paragraph, style_id: str) -> CT_Style | None:
    return block.part.package.main_document_part.styles.element.get_by_id(style_id)


def has(block: _Cell | Paragraph, style_id: str) -> bool:
    """Whether the block's document defines the style: optional parts follow the template."""
    return style(block, style_id) is not None


def section(cell: _Cell, spacing_id: str, heading_id: str, title: str) -> None:
    para(cell, spacing_id)
    para(cell, heading_id, title)


def bullet_num_id(doc: Doc) -> str:
    """numId of the template's first bullet list, so bullets match the template."""
    numbering = doc.part.numbering_part.element
    # The first bullet definition, not the first instance of any bullet definition: a template's own list is
    # often registered after python-docx's built-in ones.
    first = numbering.xpath("./w:abstractNum[w:lvl[1]/w:numFmt/@w:val='bullet']/@w:abstractNumId")[0]
    return str(numbering.xpath(f"./w:num[w:abstractNumId/@w:val='{first}']/@w:numId")[0])


def bullet(cell: _Cell, text: str, num_id: str) -> None:
    p = para(cell, "ListParagraph", text)
    p._p.get_or_add_pPr().append(xml(f'<w:numPr><w:ilvl w:val="0"/><w:numId w:val="{num_id}"/></w:numPr>'))


def hyperlink(p: Paragraph, label: str, url: str, underline: bool = True) -> None:
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), p.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = p.add_run(label)
    run.underline = underline or None
    link.append(run._r)
    p._p.append(link)


def styled(run: Run, style_id: str) -> Run:
    run._r.style = style_id
    return run


def cells(tables: Iterable[Table]) -> Iterator[_Cell]:
    """Every cell, nested ones included; a merged cell once."""
    seen = set()
    for t in tables:
        for row in t.rows:
            for c in row.cells:
                if c._tc not in seen:
                    seen.add(c._tc)
                    yield c
                    yield from cells(c.tables)


def drop(p: Paragraph) -> None:
    p._p.getparent().remove(p._p)


def photo(p: Paragraph, cv: CV) -> None:
    spec = p.text.strip("{}").split()
    for r in p.runs:
        p._p.remove(r._r)
    if "photo" in cv:
        run = p.add_run()
        run.add_picture(str(cv["dir"] / cv["photo"]), width=Cm(float(spec[1])))
        if "gray" in spec:
            run._r.xpath(".//a:blip")[0].append(OxmlElement("a:grayscl"))


def contact(p: Paragraph, cv: CV) -> None:
    for link in cv["links"]:
        q = p.insert_paragraph_before()
        q._p.style = "CVContact"
        hyperlink(q, link["label"], link["url"], underline=False)
        styled(q.add_run("  ◆"), "CVAccent")
    drop(p)


def inner_width(cell: _Cell) -> Length:
    """The cell's width less its left and right margins: what a nested table may fill."""
    pad = sum(int(w) for w in cell._tc.xpath("./w:tcPr/w:tcMar/*[self::w:left or self::w:right]/@w:w"))
    return Emu((cell.width or 0) - Twips(pad))


def highlights(p: Paragraph, cell: _Cell, cv: CV) -> None:
    items = cv.get("highlights", [])
    t = cell.add_table(rows=2, cols=len(items))
    for col, c in zip(t._tbl.tblGrid.gridCol_lst, t.columns, strict=True):
        col.w = Emu(inner_width(cell) // len(items))
        for x in c.cells:
            x.width = col.w
    for i, h in enumerate(items):
        for row, style_id, text in ((0, "CVHighlightValue", h["value"]), (1, "CVHighlightLabel", h["label"])):
            q = t.cell(row, i).paragraphs[0]
            q._p.style = style_id
            add_text(q, text)
    drop(p)


def header(doc: Doc, cv: CV) -> None:
    """Fill the template's placeholders, wherever the template puts them."""
    found: list[tuple[Paragraph, _Cell | None]] = [(p, None) for p in doc.paragraphs]
    found += [(p, c) for c in cells(doc.tables) for p in c.paragraphs]
    for p, cell in found:
        match p._p.style:
            case "Name":
                p.runs[0].text = cv["name"]
            case "JobTitle":
                p.runs[0].text = cv["title"].upper()
            case "CVPhoto" if p.text.startswith("{Photo"):
                photo(p, cv)
            case "CVContact" if p.text == "{Contact}":
                contact(p, cv)
            case "CVHighlights" if cell and p.text == "{Highlights}":
                highlights(p, cell, cv)


def fill(cell: _Cell, style_id: str) -> str:
    """The shading colour a character style defines."""
    found = style(cell, style_id)
    assert found is not None, f"template lacks {style_id}"
    return str(found.xpath(".//w:shd/@w:fill")[0])


def bar(cell: _Cell, level: int) -> None:
    """A skill level as a bar: SkillFill's shading over level/5 of the text width, SkillTrack's over the rest."""
    inner = inner_width(cell)
    parts = [(level, fill(cell, "CVSkillFill"))] + [(5 - level, fill(cell, "CVSkillTrack"))] * (level < 5)
    t = cell.add_table(rows=1, cols=len(parts))
    t._tbl.tblPr.get_or_add_tblLayout().set(qn("w:type"), "fixed")
    t.rows[0].height, t.rows[0].height_rule = Cm(0.12), WD_ROW_HEIGHT_RULE.EXACTLY
    for col, c, (share, color) in zip(t._tbl.tblGrid.gridCol_lst, t.rows[0].cells, parts, strict=True):
        col.w = c.width = Emu(inner * share // 5)
        c._tc.get_or_add_tcPr().append(shading(color))
        c.paragraphs[0]._p.style = "CVSkillBar"
    cell.paragraphs[-1]._p.style = "CVSkillBar"  # the paragraph Word needs after a table


def sidebar(cell: _Cell, cv: CV, lb: Labels) -> None:
    section(cell, "SidebarSectionsSpacing", "Heading3", lb["details"])
    para(cell, "SidebarText", cv["location"])
    para(cell, "Heading4", lb["born"].upper())
    para(cell, "SidebarText", str(cv["born"]))
    para(cell, "Heading4", lb["nationality"].upper())
    para(cell, "SidebarText", cv["nationality"])
    if not has(cell, "CVContact"):  # otherwise the header lists the links
        section(cell, "SidebarSectionsSpacing", "Heading3", lb["links"])
        for link in cv["links"]:
            hyperlink(para(cell, "SidebarText"), link["label"], link["url"])
    section(cell, "SidebarSectionsSpacing", "Heading3", lb["skills"])
    for s in cv["skills"]:
        if has(cell, "CVSkillBar"):
            para(cell, "CVSkill", s["name"])
            bar(cell, s["level"])
        else:
            para(cell, "SidebarText", f"{s['name']} · {lb['levels'][s['level']]}")
    section(cell, "SidebarSectionsSpacing", "Heading3", lb["languages"])
    for lang in cv["languages"]:
        if has(cell, "CVKeyValue"):
            styled(para(cell, "CVKeyValue", lang["name"]).add_run(f"\t{lang['level']}"), "CVMuted")
        else:
            para(cell, "SidebarText", f"{lang['name']} · {lang['level']}")
    section(cell, "SidebarSectionsSpacing", "Heading3", lb["hobbies"])
    para(cell, "SidebarText", cv["hobbies"])


def tools(cell: _Cell, names: list[str], lb: Labels) -> None:
    if not has(cell, "CVTool"):
        p = para(cell, "Normal")
        p.add_run(f"{lb['tools']}: ").italic = True
        p.add_run(", ".join(names))
        return
    p = para(cell, "CVTools")
    for name in names:
        # A shaded narrow no-break space draws the edge; a bar glyph made Word break chips mid-word.
        styled(p.add_run("\u202f"), "CVToolMark")
        styled(p.add_run(f"\u00a0{name}\u00a0"), "CVTool")
        p.add_run(" \u200b")  # Word won't break at a space next to a no-break space; a zero-width space breaks


def main_body(cell: _Cell, cv: CV, lb: Labels, num_id: str) -> None:
    cell._tc.remove(cell.paragraphs[0]._p)  # template placeholder
    para(cell, "Heading1", lb["summary"])
    for text in cv["summary"]:
        para(cell, "Normal", text)
    section(cell, "MainSectionsSpacing", "Heading1", lb["education"])
    for ed in cv["education"]:
        para(cell, "Heading2", f"{ed['degree']}, {ed['school']}, {ed['city']}")
        para(cell, "Date", span(ed, lb))
        para(cell, "Normal", ed.get("notes", ""))
    section(cell, "MainSectionsSpacing", "Heading1", lb["experience"])
    for job in cv["experience"]:
        if has(cell, "CVRole"):
            para(cell, "CVRole", job["title"])
            para(cell, "CVCompany", f"{job['company']}, {job['city']}")
        else:
            para(cell, "Heading2", f"{job['title']} {lb['at']} {job['company']}, {job['city']}")
        para(cell, "CVRoleDate" if has(cell, "CVRoleDate") else "Date", span(job, lb))
        if not has(cell, "CVTool"):
            para(cell, "Normal").add_run(lb["achievements"]).italic = True
        for text in job["bullets"]:
            bullet(cell, text, num_id)
        tools(cell, job["tools"], lb)
    section(cell, "MainSectionsSpacing", "Heading1", lb["testimonials"])
    for t in cv["testimonials"]:
        para(cell, "Heading2", f"{t['name']} | {t['role']}")
        para(cell, "Date", when(t["date"], lb).upper())
        quote = [*t["quote"]]
        quote[0] = f"“{quote[0]}"
        quote[-1] += "”"
        for text in quote:
            para(cell, "CVQuote" if has(cell, "CVQuote") else "Normal", text)


def render(cv: CV, template: Path) -> Doc:
    doc = Document(str(template))
    row = doc.tables[-1].rows[0].cells
    header(doc, cv)
    sidebar(row[0], cv, cv["labels"])
    main_body(row[-1], cv, cv["labels"], bullet_num_id(doc))
    return doc
